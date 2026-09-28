"""Regression tests for the P3 exporter review findings (1, 9, 10, 11, 12, 22, 23, 24, 25, 26)."""

from __future__ import annotations

import importlib.util
import io
import json
import os
import re
import struct
import zipfile
from pathlib import Path

import pytest

from insia_agents.exporters import ExportError, ExportFile, export_item, export_run_zip
from insia_agents.exporters import api as export_api
from insia_agents.exporters.common import clean_draft, find_placeholders, inline_segments, xml_safe
from insia_agents.exporters.instagram import (RenderedSlides, Slide, build_carousel_zip, parse_carousel, slide_warnings,
                                              slides_html)
from insia_agents.exporters.naver_html import extract_fragment, naver_fragment, naver_preview_page
from insia_agents.models import Brief, ContentItem, ContentItemDetail, Draft, DraftVersion, Profile, TeamMember

HAS_DOCX = importlib.util.find_spec("docx") is not None
HAS_PLAYWRIGHT = importlib.util.find_spec("playwright") is not None
needs_docx = pytest.mark.skipif(not HAS_DOCX, reason='python-docx 미설치 (pip install "insia-smartagent[export]")')
XML_ILLEGAL = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f]")
YELLOW = '<w:highlight w:val="yellow"/>'


def _chromium() -> str | None:
    for candidate in (os.environ.get("INSIA_CHROMIUM"), "/opt/pw-browsers/chromium"):
        if candidate and Path(candidate).exists():
            return candidate
    return None


CHROMIUM = _chromium()
needs_browser = pytest.mark.skipif(not HAS_PLAYWRIGHT or CHROMIUM is None,
                                   reason="Playwright/Chromium 없음 — 브라우저 테스트 생략")


@pytest.fixture(scope="module")
def chromium_browser():
    """One Chromium for the page-level browser tests in this module."""
    from playwright.sync_api import sync_playwright

    with sync_playwright() as pw:
        browser = pw.chromium.launch(executable_path=CHROMIUM)
        try:
            yield browser
        finally:
            browser.close()


@pytest.fixture(autouse=True)
def no_png_rendering(monkeypatch):
    monkeypatch.setenv("INSIA_RENDER", "0")
    monkeypatch.delenv("INSIA_CHROMIUM", raising=False)


def detail(channel: str, content: str, *, title: str = "테스트 제목", hashtags: list[str] | None = None,
           versions: list[DraftVersion] | None = None, item_version: int = 1) -> ContentItemDetail:
    draft = Draft(channel=channel, round=0, title=title, content=content, hashtags=hashtags or [])
    versions = versions or [DraftVersion(id="v1", item_id="it1", version=1, source="agent", draft=draft,
                                         created_at="2026-09-28T01:00:00Z")]
    item = ContentItem(id="it1", channel=channel, title=title, version=item_version, created_at="2026-09-28T01:00:00Z")
    return ContentItemDetail(item=item, versions=versions)


def unzip(data: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def document_xml(data: bytes) -> str:
    return unzip(data)["word/document.xml"].decode("utf-8")


def highlighted(xml: str, text: str) -> bool:
    """True when a run starting with ``text`` carries the yellow highlight."""
    return re.search(re.escape(YELLOW) + r"</w:rPr><w:t[^>]*>" + re.escape(text), xml) is not None


# ---------------------------------------------------------------------------
# Finding 1 — XML-illegal control characters, robust zips
# ---------------------------------------------------------------------------


def test_xml_safe_and_clean_draft_turn_soft_breaks_into_line_breaks():
    assert xml_safe("a\x0bb\x0cc\x00d\x1be\ttab\n") == "a\nb\ncd" + "e\ttab\n"
    assert xml_safe("정상 텍스트") == "정상 텍스트"
    draft = Draft(channel="bizplan", round=0, title="제목\x1b 끝", hashtags=["#태그\x00"],
                  content="## 1\n본문\x0b다음 줄\n| 단계 | 내용 |\n|---|---|\n| 1 | 개발\x0b테스트 |\n")
    cleaned = clean_draft(draft)
    assert cleaned.title == "제목 끝" and cleaned.hashtags == ["#태그"]
    assert "본문\n다음 줄" in cleaned.content
    assert "| 1 | 개발<br>테스트 |" in cleaned.content  # the table row stays one row
    assert clean_draft(cleaned) is cleaned  # nothing to clean → same object


@needs_docx
@pytest.mark.parametrize("channel", ["bizplan", "linkedin", "naver_blog", "instagram"])
@pytest.mark.parametrize(("content", "title"), [
    ("# 계획서\n## 1. 문제\n본문 (PDF에서 복사)\x0c다음 쪽", "계획서"),
    ("## 1\n본문 (PPT에서 복사)\x0b다음 줄", "t"),
    ("본문", "제목\x1b"),
    ("본문\x00 끝", "t"),
])
def test_docx_survives_xml_illegal_control_characters(channel, content, title):
    import docx

    exported = export_item(detail(channel, content, title=title, hashtags=["#a\x07"]), "docx")
    document = docx.Document(io.BytesIO(exported.data))  # a valid Word file
    text = "\n".join(p.text for p in document.paragraphs)
    assert not XML_ILLEGAL.search(text)
    if "\x0b" in content or "\x0c" in content:
        assert "다음" in text and "<w:br/>" in document_xml(exported.data)  # soft break kept as a line break


@needs_docx
def test_docx_blind_notice_with_control_character_in_profile_name():
    profile = Profile(team=[TeamMember(role="대표", name="김\x0b철수")])
    exported = export_item(detail("bizplan", "대표 김\x0b철수입니다"), "docx", profile)
    assert "블라인드 확인 필요" in document_xml(exported.data)


@needs_docx
def test_build_docx_wraps_unexpected_errors_in_korean_export_error(monkeypatch):
    from insia_agents.exporters import docx_writer

    def boom(self):
        raise RuntimeError("lxml exploded")

    monkeypatch.setattr(docx_writer._Builder, "build", boom)
    with pytest.raises(ExportError, match=r"Word\(.docx\) 파일을 만들지 못했어요 \(오류: RuntimeError\)"):
        docx_writer.build_docx(Draft(channel="bizplan", round=0, title="t", content="본문"))


def test_export_item_reports_unexpected_failures_as_export_error(monkeypatch):
    def boom(ctx):
        raise KeyError("oops")

    monkeypatch.setitem(export_api._EXPORTERS, "html", boom)
    with pytest.raises(ExportError, match="파일을 만들지 못했어요 \\(오류: KeyError\\).*다른 형식"):
        export_item(detail("naver_blog", "본문"), "html")


def test_item_zip_keeps_other_formats_when_one_fails(monkeypatch):
    def boom(ctx):
        raise ValueError("All strings must be XML compatible")

    monkeypatch.setitem(export_api._EXPORTERS, "docx", boom)
    bundle = export_item(detail("linkedin", "링크드인 본문"), "zip")
    names = unzip(bundle.data)
    assert any(n.endswith(".txt") for n in names) and any(n.endswith(".md") for n in names)
    assert not any(n.endswith(".docx") for n in names)
    assert any("docx" in n and "ValueError" in n for n in bundle.notes)


def _workspace_run(tmp_path, bizplan_content: str):
    from insia_agents.db import Workspace

    ws = Workspace(tmp_path / "ws")
    brief = Brief(topic="테스트 주제", channels=["bizplan", "linkedin"])
    run = "r_20260928_abc"
    ws.create_run(run, brief, kind="pipeline")
    ws.ensure_item(f"it_{run}_bizplan", "bizplan", "사업계획서", run_id=run, brief=brief)
    ws.add_version(f"it_{run}_bizplan", Draft(channel="bizplan", round=0, title="사업계획서", content=bizplan_content),
                   source="agent", run_id=run)
    ws.ensure_item(f"it_{run}_linkedin", "linkedin", "링크드인", run_id=run, brief=brief)
    ws.add_version(f"it_{run}_linkedin", Draft(channel="linkedin", round=0, title="링크드인", content="링크드인 본문"),
                   source="agent", run_id=run)
    ws.update_run(run, status="completed")
    return ws, run


@needs_docx
def test_run_zip_with_pasted_control_characters_end_to_end(tmp_path):
    ws, run = _workspace_run(tmp_path, "## 1. 문제\n본문 (PPT에서 복사)\x0b다음 줄\x0c다음 쪽")
    files = unzip(export_run_zip(ws, run).data)
    docx_name = next(n for n in files if n.endswith(".docx"))
    assert "bizplan/" in docx_name and files[docx_name].startswith(b"PK")
    assert any("linkedin/" in n and n.endswith(".txt") for n in files)
    item_docx = export_item(ws.get_item(f"it_{run}_bizplan"), "docx")
    assert "다음 쪽" in document_xml(item_docx.data)


def test_run_zip_never_fails_as_a_whole_when_one_format_crashes(tmp_path, monkeypatch):
    ws, run = _workspace_run(tmp_path, "## 1. 문제\n본문")
    real_export_item = export_api.export_item

    def flaky(detail_, fmt, profile=None, **kw):
        if detail_.item.channel == "bizplan":
            raise RuntimeError("unexpected")
        return real_export_item(detail_, fmt, profile, **kw)

    monkeypatch.setattr(export_api, "export_item", flaky)
    exported = export_run_zip(ws, run)
    files = unzip(exported.data)
    assert any("linkedin/" in n and n.endswith(".txt") for n in files)
    assert not any("bizplan/" in n and not n.endswith("review.json") for n in files)
    assert any(n.startswith("사업계획서 docx: 파일을 만들지 못했어요 (오류: RuntimeError)") for n in exported.notes)
    readme = next(v for k, v in files.items() if k.endswith("README.txt")).decode("utf-8")
    assert "파일을 만들지 못했어요" in readme


def test_naver_html_and_txt_drop_control_characters():
    content = "## 소제목\n본문\x0b다음 줄\x00 끝\n\n#a #b"
    item = detail("naver_blog", content, title="제목\x1b", hashtags=["#a", "#b"])
    page = export_item(item, "html").data.decode("utf-8")
    assert not XML_ILLEGAL.search(page)
    assert "본문<br>다음 줄 끝" in extract_fragment(page)
    text = export_item(item, "txt").data.decode("utf-8")
    assert not XML_ILLEGAL.search(text) and "본문\n다음 줄 끝" in text and text.startswith("제목\n")
    assert not XML_ILLEGAL.search(export_item(detail("linkedin", "첫 줄\x0b둘째 줄"), "txt").data.decode("utf-8"))


# ---------------------------------------------------------------------------
# Finding 9 — Instagram labels
# ---------------------------------------------------------------------------

CAPTION = "\n\n## 캡션\n첫 줄 훅이에요\n저장해 두세요\n#a #b #c\n"


def _slide(body: str, n: int = 1) -> str:
    return f"## 캐러셀\n### 슬라이드 {n} — 훅\n{body}\n"


@pytest.mark.parametrize("body", [
    "- **문구:** 슬라이드 문구\n- **비주얼:** 남색 배경\n- **대체텍스트:** 대체",
    "- **문구**: 슬라이드 문구\n- **비주얼**: 남색 배경\n- **대체텍스트**: 대체",
    "- 헤드라인: 슬라이드 문구\n- 비주얼 지시: 남색 배경\n- 대체 텍스트(alt): 대체",
    "**문구:** 슬라이드 문구\n비주얼: 남색 배경\n대체텍스트: 대체",
    "- **문구: 슬라이드 문구**\n- 디자인: 남색 배경\n- ALT: 대체",
])
def test_parse_carousel_accepts_bold_and_alias_labels(body):
    slide = parse_carousel(_slide(body))[0]
    assert slide.parts == ["슬라이드 문구"]
    assert slide.visual == "남색 배경" and slide.alt == "대체" and slide.skipped == ()


def test_unknown_labels_never_reach_the_image():
    # "레이아웃 노트" is a design direction: it goes to 비주얼 (README), not to the skipped list.
    body = ("- 문구: 좋은 문구\n- 포인트: 좌측 정렬, 여백 크게\n  두 번째 줄 메모\n- 대체텍스트: 대체\n"
            "- 톤: 밝게\n- 레이아웃 노트: 여백 크게")
    slide = parse_carousel(_slide(body))[0]
    assert slide.parts == ["좋은 문구"] and slide.alt == "대체"
    assert slide.skipped == ("포인트", "톤") and slide.visual == "여백 크게"
    page = slides_html([slide])
    assert "좌측 정렬" not in page and "두 번째 줄 메모" not in page and "밝게" not in page
    assert "여백 크게" not in page


def test_content_lines_with_colons_stay_on_the_image():
    body = ("- 문구: 도입 전 체크리스트\n  - 목표: 무엇을 얻을지\n- 1단계: 자료 모으기\n- 오후 2:00 시작\n"
            "- 서브: 저장해 두세요\n- 비주얼: 체크 아이콘\n- 대체텍스트: 체크리스트 이미지")
    slide = parse_carousel(_slide(body))[0]
    assert slide.parts == ["도입 전 체크리스트", "목표: 무엇을 얻을지", "1단계: 자료 모으기", "오후 2:00 시작",
                           "저장해 두세요"]
    assert slide.skipped == ()


def test_slide_text_source_and_alt_never_show_markdown_bold():
    slide = parse_carousel(_slide("- 문구: 핵심 **3가지** 정리 / **저장**하세요\n- 대체텍스트: **굵은** 설명\n"
                                  "- 출처: **통계청** 2025"))[0]
    assert slide.parts == ["핵심 3가지 정리", "저장하세요"]
    page = slides_html([slide])
    assert "**" not in page and "출처: 통계청 2025" in page
    files = unzip(build_carousel_zip(Draft(channel="instagram", round=0, title="t",
                                           content=_slide("- 문구: a\n- 대체텍스트: **굵은** 설명") + CAPTION))[0])
    assert "굵은 설명" in files["alt-text.txt"].decode("utf-8") and "**" not in files["alt-text.txt"].decode("utf-8")


def test_carousel_zip_notes_missing_text_alt_and_skipped_labels():
    content = ("## 캐러셀\n### 슬라이드 1 — 훅\n- **문구:** 첫 장\n- **대체텍스트:** 대체1\n"
               "### 슬라이드 2 — 문제\n- 비주얼: 남색\n- 대체텍스트: 대체2\n"
               "### 슬라이드 3 — 끝\n- 문구: 끝 / 저장\n- 메모: 내부용\n" + CAPTION)
    data, notes = build_carousel_zip(Draft(channel="instagram", round=0, title="t", content=content))
    joined = "\n".join(notes)
    assert "문구(- 문구:)를 찾지 못한 슬라이드: 2" in joined
    assert "대체텍스트(- 대체텍스트:)를 찾지 못한 슬라이드: 3" in joined
    assert "슬라이드 3(메모)" in joined
    files = unzip(data)
    readme = files["README.txt"].decode("utf-8")
    assert "먼저 확인할 부분" in readme and "찾지 못한 슬라이드: 2" in readme
    alt = files["alt-text.txt"].decode("utf-8")
    assert "대체1" in alt and "대체2" in alt
    assert "내부용" not in files["slides.html"].decode("utf-8")


def test_sample_style_carousel_has_no_warnings():
    content = _slide("- 문구: 좋은 문구 / 보조 문구\n- 비주얼: 남색\n- 대체텍스트: 대체\n- 출처: 통계청") + CAPTION
    _, notes = build_carousel_zip(Draft(channel="instagram", round=0, title="t", content=content))
    assert notes == ("PNG 대신 slides.html을 넣었어요 — PNG 렌더링이 꺼져 있어요 (INSIA_RENDER=0)",)


@pytest.mark.parametrize("body", [
    "문구: 좋은 문구\n비주얼 지시사항: 남색 배경, 큰 흰 글씨\n대체텍스트: 대체",  # unbulleted
    "**문구:** 좋은 문구\n비주얼 지시사항: 남색 배경, 큰 흰 글씨\n**대체텍스트:** 대체",  # bold + plain mix
    "- 문구: 좋은 문구\n  비주얼 지시사항: 남색 배경, 큰 흰 글씨\n- 대체텍스트: 대체",  # indented under 문구
    "- 문구: 좋은 문구\n비주얼 지시사항: 남색 배경, 큰 흰 글씨\n- 대체텍스트: 대체",  # unbulleted under a bullet
    "1. 문구: 좋은 문구\n2. 비주얼: 남색 배경, 큰 흰 글씨\n3. 대체텍스트: 대체",  # numbered labels
    "1) 문구: 좋은 문구\n2) 디자인 디렉션: 남색 배경, 큰 흰 글씨\n3) 대체 텍스트: 대체",
    "- 📌 문구: 좋은 문구\n- 🎨 비주얼: 남색 배경, 큰 흰 글씨\n- ♿ 대체텍스트: 대체",  # emoji before the label
    "- 문구: 좋은 문구\n- Image: 남색 배경, 큰 흰 글씨\n- Alt text: 대체",
])
def test_design_directions_go_to_visual_in_every_label_style(body):
    slide = parse_carousel(_slide(body))[0]
    assert slide.parts == ["좋은 문구"]
    assert slide.visual == "남색 배경, 큰 흰 글씨" and slide.alt == "대체" and slide.skipped == ()
    assert "남색 배경" not in slides_html([slide])


@pytest.mark.parametrize(("body", "skipped"), [
    ("문구: 좋은 문구\n메모: 디자이너 전달용 내부 메모\n대체텍스트: 대체", ("메모",)),
    ("- 문구: 좋은 문구\n메모: 디자이너 전달용 내부 메모\n- 대체텍스트: 대체", ("메모",)),
    ("- 문구: 좋은 문구\n  - 메모: 디자이너 전달용 내부 메모\n- 대체텍스트: 대체", ("메모",)),
    ("- 문구: 좋은 문구\n  디자이너 메모: 디자이너 전달용 내부 메모\n- 대체텍스트: 대체", ("디자이너메모",)),
    ("1. 문구: 좋은 문구\n2. 메모: 디자이너 전달용 내부 메모\n3. 포인트: 강조\n4. 대체텍스트: 대체", ("메모", "포인트")),
    ("문구: 좋은 문구\n포인트: 디자이너 전달용 내부 메모\n대체텍스트: 대체", ("포인트",)),  # unknown sibling label
    ("**문구:** 좋은 문구\n**제작 노트:** 디자이너 전달용 내부 메모\n**대체텍스트:** 대체", ("제작노트",)),
    ("  - 문구: 좋은 문구\n  - 메모: 디자이너 전달용 내부 메모\n  - 대체텍스트: 대체", ("메모",)),  # whole body indented
])
def test_memos_and_unknown_labels_stay_off_the_image_in_every_label_style(body, skipped):
    slide = parse_carousel(_slide(body))[0]
    assert slide.parts == ["좋은 문구"] and slide.alt == "대체" and slide.skipped == skipped
    assert "내부 메모" not in slides_html([slide])
    _, notes = build_carousel_zip(Draft(channel="instagram", round=0, title="t", content=_slide(body) + CAPTION))
    assert any(f"슬라이드 1({', '.join(skipped)})" in note and "이미지에 넣지 않은 항목" in note for note in notes)


def test_content_lines_under_the_copy_stay_on_the_image_in_every_label_style():
    # a list under an unbulleted 문구, and nested or numbered lines under a bulleted 문구, are content
    plain = parse_carousel(_slide("문구: 창업 전 체크리스트\n- 자금: 3개월 운영비\n- 인력: 1인\n대체텍스트: 대체"))[0]
    assert plain.parts == ["창업 전 체크리스트", "자금: 3개월 운영비", "인력: 1인"] and plain.skipped == ()
    nested = parse_carousel(_slide("- 문구: 세 가지 습관\n  1. 기록: 매일 적기\n  2. 검토: 주 1회\n"
                                   "  이유: 쌓이면 보여요\n- 대체텍스트: 대체"))[0]
    assert nested.parts == ["세 가지 습관", "1. 기록: 매일 적기", "2. 검토: 주 1회", "이유: 쌓이면 보여요"]
    assert nested.skipped == ()
    numbered = parse_carousel(_slide("1. 문구: 세 가지 습관\n- 기록: 매일 적기\n2. 대체텍스트: 대체"))[0]
    assert numbered.parts == ["세 가지 습관", "기록: 매일 적기"] and numbered.alt == "대체"
    no_labels = parse_carousel(_slide("1. 기록: 매일 적기\n2. 검토: 주 1회"))[0]  # a slide that is only content
    assert no_labels.parts == ["1. 기록: 매일 적기", "2. 검토: 주 1회"] and no_labels.skipped == ()


def test_title_and_body_labels_fill_a_slide_without_copy():
    slide = parse_carousel(_slide("- 제목: 쓰는 시간보다\n- 본문: 고치는 시간이 길어요\n- 대체텍스트: 대체"))[0]
    assert slide.parts == ["쓰는 시간보다", "고치는 시간이 길어요"] and slide.skipped == ()
    both = parse_carousel(_slide("- 제목: 훅\n- 문구: 쓰는 시간보다\n- 대체텍스트: 대체"))[0]
    assert both.parts == ["쓰는 시간보다"] and both.skipped == ("제목",)  # 문구 wins; 제목 is reported


# ---------------------------------------------------------------------------
# Finding 10 — what counts as a placeholder
# ---------------------------------------------------------------------------


def test_placeholders_skip_disclosures_citation_lists_and_checkboxes():
    text = ("[광고] 이 글은 자사 서비스 소개 글이에요.\n시장 규모 3조 원 [s2, s5] [s2·s5] [s3~s5]\n[유료 광고 포함]\n"
            "[확인 필요: 2025 수치]\n[TIP] 팁\n[협찬]\n[AD]\n- [ ] 체크\n- [x] 완료\n[○]명 [광고 문구: 여기에]\n"
            "[광고주명] [대표자 성명] [s12] [별첨 1]")
    assert find_placeholders(text) == ["[확인 필요: 2025 수치]", "[○]", "[광고 문구: 여기에]", "[광고주명]", "[대표자 성명]"]
    assert find_placeholders(text, keep=["[대표자 성명] 드림"]) == \
        ["[확인 필요: 2025 수치]", "[○]", "[광고 문구: 여기에]", "[광고주명]"]


def test_naver_preview_never_asks_to_delete_a_required_disclosure():
    draft = Draft(channel="naver_blog", round=0, title="t",
                  content="[이 글은 협찬 광고예요] 본문\n\n## 소제목\n시장 [s2, s5]\n\n- [ ] 체크\n\n[확인 필요: 수치]")
    profile = Profile(required_phrases=["[이 글은 협찬 광고예요]"])
    page = naver_preview_page(draft, profile=profile)
    panel = re.search(r"<h2>게시 전에 채우거나 지워야 할 자리 (\d+)곳</h2><ul>(.*?)</ul>", page)
    assert panel.group(1) == "1" and panel.group(2) == "<li>[확인 필요: 수치]</li>"
    fragment = extract_fragment(page)
    assert fragment == naver_fragment(draft, profile)
    assert '<span style="background-color:#FFF3B0;">[이 글은 협찬 광고예요]</span>' not in fragment
    assert "[이 글은 협찬 광고예요] 본문" in fragment
    assert '<span style="background-color:#FFF3B0;">[확인 필요: 수치]</span>' in fragment
    assert "#FFF3B0;\">[s2, s5]" not in fragment and "#FFF3B0;\">[ ]" not in fragment


@pytest.mark.parametrize("disclosure", [
    "[광고]", "[협찬]", "[유료 광고 포함]", "[유료광고]", "[광고·협찬]", "[광고・협찬]", "[광고/협찬]", "[광고, 협찬]",
    "[광고 및 협찬]", "[제휴 광고]", "[체험단]", "[AD]", "[Ad]", "[Sponsored]", "[Paid partnership]", "[ 광고 ]",
])
def test_literal_ad_disclosures_are_not_placeholders(disclosure):
    assert find_placeholders(f"{disclosure} 이 글은 자사 서비스 소개예요") == []


@pytest.mark.parametrize("fill_in", [
    "[광고 예산]", "[광고 대행사 이름]", "[체험단 모집 인원]", "[스폰서 이름]", "[AD campaign]", "[광고 표시]",
    "[광고 예산 확인 필요]", "[광고비 ○○만 원]", "[협찬사 이름]", "[제휴사명]", "[광고 포함 여부]",
])
def test_fill_ins_that_start_with_a_disclosure_word_stay_placeholders(fill_in):
    assert find_placeholders(f"□ (마케팅) 온라인 광고 {fill_in} 집행") == [fill_in]


def test_disclosure_word_fill_in_is_listed_and_highlighted_everywhere():
    content = "## 마케팅\n[광고] 이 글은 자사 서비스 소개예요\n\n월 광고비 [광고 예산 확인 필요]"
    draft = Draft(channel="naver_blog", round=0, title="t", content=content)
    page = naver_preview_page(draft)
    panel = re.search(r"<h2>게시 전에 채우거나 지워야 할 자리 (\d+)곳</h2><ul>(.*?)</ul>", page)
    assert panel.group(1) == "1" and panel.group(2) == "<li>[광고 예산 확인 필요]</li>"
    fragment = naver_fragment(draft)
    assert '<span style="background-color:#FFF3B0;">[광고 예산 확인 필요]</span>' in fragment
    assert '#FFF3B0;">[광고]' not in fragment
    if HAS_DOCX:
        xml = document_xml(export_item(detail("bizplan", "## 3. 성장전략\n□ (마케팅) 월 광고비 [광고 예산 확인 필요]\n"
                                                         "\n[광고]"), "docx").data)
        assert highlighted(xml, "[광고 예산 확인 필요]") and not highlighted(xml, "[광고]")


@needs_docx
def test_docx_highlights_only_real_fill_ins():
    content = "## 1. 시장\n시장 규모 3조 원 [s2, s5]\n\n[필수 고지 문구]\n\n□ (고객) [확인 필요: 고객 수]"
    xml = document_xml(export_item(detail("bizplan", content), "docx",
                                   Profile(required_phrases=["[필수 고지 문구]"])).data)
    assert not highlighted(xml, "[s2, s5]") and not highlighted(xml, "[필수 고지 문구]")
    assert highlighted(xml, "[확인 필요: 고객 수]")


# ---------------------------------------------------------------------------
# Finding 11 — reference headings
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(("heading", "expected"), [
    ("참고자료", True), ("5. 참고 자료", True), ("참고문헌", True), ("출처", True), ("출처 목록", True),
    ("Ⅴ. 참고자료 및 출처", True), ("References", True), ("**참고자료**", True),
    ("3-2. 국내 매출처 확보 계획", False), ("2-1. 학습 데이터 출처와 수집 방법", False), ("출처 표기 원칙", False),
    ("참고로 알아 둘 점", False), ("", False),
])
def test_reference_heading_is_matched_as_a_whole(heading, expected):
    from insia_agents.exporters.docx_writer import is_reference_heading

    assert is_reference_heading(heading) is expected


@needs_docx
def test_docx_placeholders_highlighted_under_sales_and_data_source_headings():
    import docx

    content = ("## 3. 성장전략\n### 3-2. 국내 매출처 확보 계획\n□ (판로) 초기 고객 [확인 필요: 파일럿 업체 수]\n"
               "### 2-1. 학습 데이터 출처와 수집 방법\n□ (데이터) 공공데이터 [확인 필요: 데이터셋 이름]\n"
               "## 참고자료\n- [s1] 통계청 [확인 필요: 발행일]\n")
    exported = export_item(detail("bizplan", content), "docx")
    xml = document_xml(exported.data)
    for placeholder in ("[확인 필요: 파일럿 업체 수]", "[확인 필요: 데이터셋 이름]", "[확인 필요: 발행일]"):
        assert highlighted(xml, placeholder), placeholder
    paragraphs = docx.Document(io.BytesIO(exported.data)).paragraphs
    sales = next(p for p in paragraphs if p.text.startswith("□ (판로)"))
    assert sales.runs[-1].font.size is None  # body text size, not the 9pt reference style
    reference = next(p for p in paragraphs if p.text.startswith("- [s1]"))
    assert reference.runs[-1].font.size.pt == 9  # the real reference list keeps its style


# ---------------------------------------------------------------------------
# Finding 12 — version in the file name, no silent overwrite
# ---------------------------------------------------------------------------


def _two_versions() -> ContentItemDetail:
    v1 = Draft(channel="linkedin", round=0, title="같은 제목", content="v1 본문")
    v2 = Draft(channel="linkedin", round=1, title="같은 제목", content="v2 본문")
    versions = [DraftVersion(id="v1", item_id="it1", version=1, source="agent", draft=v1, created_at="2026-09-28T01:00:00Z"),
                DraftVersion(id="v2", item_id="it1", version=2, source="agent", draft=v2, created_at="2026-09-28T02:00:00Z")]
    return detail("linkedin", "", title="같은 제목", versions=versions, item_version=2)


def test_export_filename_carries_the_version(tmp_path):
    item = _two_versions()
    current, old = export_item(item, "txt"), export_item(item, "txt", version=1)
    assert current.filename == "2026-09-28_linkedin_같은-제목_v2.txt"
    assert old.filename == "2026-09-28_linkedin_같은-제목_v1.txt"
    header = current.content_disposition()
    header.encode("latin-1")
    assert 'filename="2026-09-28_linkedin_v2.txt"' in header and "_v2.txt" in header.split("filename*=")[1]
    path2 = current.save(tmp_path)
    path1 = old.save(tmp_path)
    assert path1 != path2 and path2.read_text("utf-8").startswith("v2 본문")
    assert path1.read_text("utf-8").startswith("v1 본문")
    names = set(unzip(export_item(item, "zip").data))
    assert names and all("_v2." in n for n in names)


def test_save_never_overwrites_a_different_file(tmp_path):
    first = ExportFile("a_v2.txt", "text/plain; charset=utf-8", "내용".encode("utf-8"))
    path = first.save(tmp_path)
    assert first.save(tmp_path) == path  # identical bytes → same file, no duplicate
    path.write_text("내용 + 대표님이 직접 고친 부분", encoding="utf-8")
    second = first.save(tmp_path)
    assert second.name == "a_v2 (2).txt" and second.read_text("utf-8") == "내용"
    assert "직접 고친" in path.read_text("utf-8")
    assert ExportFile("a_v2.txt", "text/plain", b"other").save(tmp_path).name == "a_v2 (3).txt"
    replaced = ExportFile("a_v2.txt", "text/plain", b"new").save(tmp_path, overwrite=True)
    assert replaced == path and path.read_bytes() == b"new"


# ---------------------------------------------------------------------------
# Finding 22 — URLs followed by Korean particles
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(("text", "url"), [
    ("신청은 https://www.k-startup.go.kr에서 하세요", "https://www.k-startup.go.kr"),
    ("출처: 통계청(https://kosis.kr)，2025년", "https://kosis.kr"),
    ("… https://example.com/a。", "https://example.com/a"),
    ("링크 https://example.com/path를 보세요", "https://example.com/path"),
    ("https://ko.wikipedia.org/wiki/소상공인 문서", "https://ko.wikipedia.org/wiki/소상공인"),
    ("검색 https://search.naver.com/search.naver?query=소상공인 결과", "https://search.naver.com/search.naver?query=소상공인"),
    ("https://ex.com/a_(b).", "https://ex.com/a_(b)"),
])
def test_url_detection_stops_at_korean_particles_and_fullwidth_punctuation(text, url):
    segments = list(inline_segments(text))
    assert [s.text for s in segments if s.kind == "url"] == [url]
    assert "".join(s.text for s in segments) == text  # nothing is lost from the visible text


@pytest.mark.parametrize(("text", "url"), [
    # browsers copy root URLs with a trailing slash
    ("공고는 https://www.k-startup.go.kr/에서 확인해요", "https://www.k-startup.go.kr/"),
    ("https://www.bizinfo.go.kr/을 참고하세요", "https://www.bizinfo.go.kr/"),
    ("https://a.com/path/에서", "https://a.com/path/"),
    ("공식 사이트는 https://www.k-startup.go.kr/입니다.", "https://www.k-startup.go.kr/"),
    ("(https://kosis.kr/에서)", "https://kosis.kr/"),
    # a particle glued to a Hangul path segment at the end of the link
    ("자세한 내용은 https://ko.wikipedia.org/wiki/소상공인을 보세요", "https://ko.wikipedia.org/wiki/소상공인"),
    ("https://ko.wikipedia.org/wiki/소상공인에서 봤어요", "https://ko.wikipedia.org/wiki/소상공인"),
    # nouns that merely end like a particle stay whole
    ("https://ko.wikipedia.org/wiki/고양이 문서", "https://ko.wikipedia.org/wiki/고양이"),
    ("https://ko.wikipedia.org/wiki/국가 문서", "https://ko.wikipedia.org/wiki/국가"),
    ("https://ko.wikipedia.org/wiki/마을 문서", "https://ko.wikipedia.org/wiki/마을"),
    ("https://a.com/이벤트 참고", "https://a.com/이벤트"),
    ("https://a.com/에서/b 참고", "https://a.com/에서/b"),
    # internationalized domain names are linked again
    ("http://한국.kr/abc 참고", "http://한국.kr/abc"),
    ("https://한국.kr에서 확인", "https://한국.kr"),
    ("http://도메인.한국에서 확인", "http://도메인.한국"),
    # a bracket after the link is not part of it
    ("https://kosis.kr[확인 필요: 기준연도]", "https://kosis.kr"),
    ("http://[::1]:8080/x 보기", "http://[::1]:8080/x"),
])
def test_url_detection_after_a_trailing_slash_idn_and_brackets(text, url):
    segments = list(inline_segments(text))
    assert [s.text for s in segments if s.kind == "url"] == [url]
    assert "".join(s.text for s in segments) == text


def test_placeholder_glued_to_a_url_is_still_found():
    assert find_placeholders("출처: https://kosis.kr[확인 필요: 기준연도]") == ["[확인 필요: 기준연도]"]


def test_naver_and_docx_links_after_a_trailing_slash():
    content = "공고는 https://www.k-startup.go.kr/에서 확인해요\n\nhttps://www.bizinfo.go.kr/을 참고하세요"
    fragment = naver_fragment(Draft(channel="naver_blog", round=0, title="t", content=content))
    assert re.findall(r'href="([^"]+)"', fragment) == ["https://www.k-startup.go.kr/", "https://www.bizinfo.go.kr/"]
    assert "https://www.k-startup.go.kr/</a>에서 확인해요" in fragment
    if HAS_DOCX:
        rels = unzip(export_item(detail("naver_blog", content), "docx").data)["word/_rels/document.xml.rels"].decode()
        assert sorted(re.findall(r'Target="(https?://[^"]+)"', rels)) == \
            ["https://www.bizinfo.go.kr/", "https://www.k-startup.go.kr/"]


def test_naver_and_docx_links_do_not_include_particles():
    content = "공고는 https://www.k-startup.go.kr에서 확인해요"
    fragment = naver_fragment(Draft(channel="naver_blog", round=0, title="t", content=content))
    assert re.findall(r'href="([^"]+)"', fragment) == ["https://www.k-startup.go.kr"]
    assert "https://www.k-startup.go.kr</a>에서 확인해요" in fragment
    if HAS_DOCX:
        rels = unzip(export_item(detail("naver_blog", content), "docx").data)["word/_rels/document.xml.rels"].decode()
        assert re.findall(r'Target="(https?://[^"]+)"', rels) == ["https://www.k-startup.go.kr"]


# ---------------------------------------------------------------------------
# Finding 23 — JSON data island in the Naver preview page
# ---------------------------------------------------------------------------

TRICKY_TITLE = "AI 팁 <!--<script> 예시 & </script>"


def test_naver_preview_data_json_escapes_every_angle_bracket():
    page = naver_preview_page(Draft(channel="naver_blog", round=0, title=TRICKY_TITLE, content="## 소제목\n본문",
                                    hashtags=["#<!--<script>"]))
    raw = re.search(r'<script type="application/json" id="insia-data">(.*?)</script>', page, re.S).group(1)
    assert "<" not in raw and ">" not in raw and "&" not in raw
    data = json.loads(raw)
    assert data == {"title": TRICKY_TITLE, "tags": "#<!--<script>"}
    assert page.count("<script") == 2


@needs_browser
def test_naver_preview_copy_buttons_work_with_tricky_title(chromium_browser):
    page_html = naver_preview_page(Draft(channel="naver_blog", round=0, title=TRICKY_TITLE, content="## 소제목\n본문",
                                         hashtags=["#<!--<script>"]))
    page = chromium_browser.new_page()
    try:
        errors: list[str] = []
        page.on("pageerror", lambda exc: errors.append(str(exc)))
        page.set_content(page_html)
        assert page.evaluate("document.scripts.length") == 2
        page.click("#copy-title")
        assert "제목을 복사했어요" in page.evaluate("document.getElementById('status').textContent")
        assert errors == []
    finally:
        page.close()


# ---------------------------------------------------------------------------
# Finding 24 — blind box only for business plans
# ---------------------------------------------------------------------------


@needs_docx
@pytest.mark.parametrize(("channel", "boxed"), [("linkedin", False), ("naver_blog", False), ("instagram", False),
                                                  ("bizplan", True)])
def test_blind_notice_only_in_bizplan_docx(channel, boxed):
    profile = Profile(team=[TeamMember(role="대표", name="김철수")])
    exported = export_item(detail(channel, "INSIA 대표 김철수입니다"), "docx", profile)
    assert ("블라인드 확인 필요" in document_xml(exported.data)) is boxed
    assert bool(exported.notes) is boxed


# ---------------------------------------------------------------------------
# Finding 25 — carousel text that does not fit
# ---------------------------------------------------------------------------


def test_slide_warnings_for_overflow_and_long_text_without_browser():
    slides = [Slide(1, "a", "첫 장", alt="a"), Slide(2, "b", " ".join(["소상공인이"] * 90), alt="b"),
              Slide(3, "c", "끝", alt="c", source=" ".join(["중소벤처기업부"] * 130))]
    assert slide_warnings(slides, overflow=(2,)) == ["글이 길어 잘린 슬라이드: 2 — 문구나 출처를 줄인 뒤 다시 내보내 주세요."]
    guess = slide_warnings(slides, rendered=False)
    assert guess and guess[0].startswith("글이 길어 잘릴 수 있는 슬라이드: 2, 3")
    assert slide_warnings(slides) == []  # rendered and nothing overflowed


def test_carousel_zip_reports_rendered_overflow(monkeypatch):
    from insia_agents.exporters import instagram

    fake_png = b"\x89PNG\r\n\x1a\n" + struct.pack(">I", 13) + b"IHDR" + struct.pack(">II", 1080, 1350)

    def fake_render(page_html, count):
        rendered = RenderedSlides([fake_png] * count)
        rendered.overflow = (2,)
        return rendered

    monkeypatch.setattr(instagram, "render_pngs", fake_render)
    content = ("## 캐러셀\n" + "".join(f"### 슬라이드 {i} — t\n- 문구: 문구 {i}\n- 대체텍스트: 대체 {i}\n" for i in (1, 2, 3))
               + CAPTION)
    data, notes = build_carousel_zip(Draft(channel="instagram", round=0, title="t", content=content))
    assert notes == ("글이 길어 잘린 슬라이드: 2 — 문구나 출처를 줄인 뒤 다시 내보내 주세요.",)
    files = unzip(data)
    assert "slide-03.png" in files and "잘린 슬라이드: 2" in files["README.txt"].decode("utf-8")


def test_slides_html_marks_overflow_in_the_preview_toolbar():
    page = slides_html([Slide(1, "a", "첫 장")])
    assert '<span class="fit-warning" role="status"></span>' in page and "window.__insiaOverflow" in page
    assert "justify-content: safe center" in page
    assert "fit-warning" not in slides_html([Slide(1, "a", "첫 장")], render_mode=True).split("<body")[1].split("<script>")[0]


@needs_browser
def test_real_render_detects_overflow_and_keeps_normal_slides_clean(monkeypatch):
    monkeypatch.delenv("INSIA_RENDER", raising=False)
    monkeypatch.setenv("INSIA_CHROMIUM", CHROMIUM)
    huge = f"{' '.join(['소상공인이'] * 40)} / {' '.join(['마케팅을'] * 45)} / {' '.join(['마케팅을'] * 45)}"
    fits = f"{' '.join(['소상공인이'] * 30)} / {' '.join(['마케팅을'] * 25)} / {' '.join(['마케팅을'] * 25)}"
    content = (f"## 캐러셀\n### 슬라이드 1 — a\n- 문구: {fits}\n- 대체텍스트: a\n- 출처: 통계청\n"
               f"### 슬라이드 2 — b\n- 문구: {huge}\n- 대체텍스트: b\n- 출처: {' '.join(['중소벤처기업부'] * 15)}\n"
               "### 슬라이드 3 — c\n- 문구: 끝 / 저장하세요\n- 대체텍스트: c\n" + CAPTION)
    data, notes = build_carousel_zip(Draft(channel="instagram", round=0, title="t", content=content))
    assert notes == ("글이 길어 잘린 슬라이드: 2 — 문구나 출처를 줄인 뒤 다시 내보내 주세요.",)
    files = unzip(data)
    assert sorted(n for n in files if n.endswith(".png")) == [f"slide-{i:02d}.png" for i in range(1, 4)]
    assert struct.unpack(">II", files["slide-02.png"][16:24]) == (1080, 1350)


@needs_browser
def test_fit_script_is_idempotent_and_returns_overflowing_slides(chromium_browser):
    huge = Slide(2, "b", " ".join(["소상공인이"] * 40) + " / " + " / ".join(" ".join(["마케팅을"] * 55) for _ in range(2)))
    page = chromium_browser.new_page(viewport={"width": 1080, "height": 1350})
    try:
        page.set_content(slides_html([Slide(1, "a", "첫 장"), huge, Slide(3, "c", "끝")], render_mode=True))
        first = page.evaluate("() => window.__insiaFit()")
        assert first == [2] and page.evaluate("() => window.__insiaFit()") == first
        top = page.evaluate("() => document.querySelectorAll('.slide')[1].querySelector('.bar').getBoundingClientRect().top")
        assert top >= 84  # overflowing text no longer spills up over the counter row
    finally:
        page.close()


# ---------------------------------------------------------------------------
# Finding 26 — <br> in the paste text
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("channel", ["bizplan", "naver_blog"])
def test_txt_converts_br_like_docx_and_html(channel):
    content = ("| 단계 | 내용 |\n|---|---|\n| 1단계 | 개발<br>테스트 |\n\n문단 첫 줄<br>둘째 줄\n\n"
               "- 항목<BR/>다음\n\n> 인용<br />끝")
    text = export_item(detail(channel, content), "txt").data.decode("utf-8")
    assert "<br" not in text.lower()
    assert "1단계\t개발 / 테스트" in text and "문단 첫 줄\n둘째 줄" in text
    assert "- 항목\n  다음" in text and "인용\n끝" in text


def test_linkedin_and_instagram_txt_convert_br():
    assert export_item(detail("linkedin", "첫 줄<br>둘째 줄"), "txt").data.decode("utf-8") == "첫 줄\n둘째 줄\n"
    caption = export_item(detail("instagram", _slide("- 문구: a") + "\n## 캡션\n첫 줄<br>둘째 줄\n#a #b #c"), "txt")
    assert caption.data.decode("utf-8").startswith("첫 줄\n둘째 줄\n")


# ---------------------------------------------------------------------------
# Blind box also warns about school/employer names from the team backgrounds
# ---------------------------------------------------------------------------


@needs_docx
def test_bizplan_docx_blind_box_warns_about_schools_and_employers():
    import docx

    profile = Profile(company_name="인시아", team=[TeamMember(role="대표", name="김철수",
                                                         background="카카오 출신 PM 7년, 고려대학교 경영학 졸업")])

    def box_text(content: str) -> str:
        exported = export_item(detail("bizplan", content), "docx", profile)
        document = docx.Document(io.BytesIO(exported.data))
        return "\n".join(c.text for t in document.tables for row in t.rows for c in row.cells)

    leaked = box_text("## 4. 팀 구성\n- 대표: 카카오 출신 PM 7년, 고려대학교 경영학 졸업 (자사 자료)")
    assert "블라인드 확인 필요 — 팀 배경의 학교·직장명 2개가 본문에 보여요: 고려대학교, 카카오." in leaked
    assert "학교명·직장명을 쓸 수 없으니" in leaked
    both = box_text("대표 김철수는 카카오 출신 PM이에요.")
    assert "팀원 실명 1개가 본문에 보여요: 김철수. 팀 배경의 학교·직장명 1개도 보여요: 카카오." in both
    assert "실명·학교명·직장명을" in both
    # masked text and platform mentions are fine: no box (same rule as the reviewer's blind check)
    assert "블라인드" not in box_text("- 대표: ○○ 출신 PM 7년, ○○대학교 졸업\n- 인시아는 카카오톡 채널과 네이버 블로그로 알려요")


def test_blind_warning_text_is_shared_by_the_word_box_and_the_export_notes():
    """``docx_writer.blind_warning``: one wording for the Word red box and (via the export API) the X-Insia-Notes line."""
    from insia_agents.exporters.docx_writer import blind_warning

    profile = Profile(company_name="인시아", team=[TeamMember(role="대표", name="김철수",
                                                         background="카카오 출신 PM 7년, 고려대학교 경영학 졸업")])
    draft = detail("bizplan", "- 대표: 카카오 출신 PM 7년, 고려대학교 경영학 졸업 (자사 자료)").versions[-1].draft
    assert blind_warning(draft, profile) == ("블라인드 확인 필요 — 팀 배경의 학교·직장명 2개가 본문에 보여요: 고려대학교, 카카오. "
                                             "사업계획서 제출본에는 학교명·직장명을 쓸 수 없으니 ○○로 가려 주세요.")
    named = draft.model_copy(update={"content": "대표 김철수가 만들어요."})
    assert blind_warning(named, profile).startswith("블라인드 확인 필요 — 팀원 실명 1개가 본문에 보여요: 김철수.")
    assert blind_warning(named.model_copy(update={"content": "대표 ○○가 만들어요."}), profile) == ""
    assert blind_warning(draft, None) == ""


def test_bizplan_docx_export_notes_carry_the_word_box_warning():
    """The docx export's notes (X-Insia-Notes) name school/employer leaks too, in the Word box's own words."""
    pytest.importorskip("docx")
    from insia_agents.exporters.docx_writer import blind_warning

    profile = Profile(company_name="인시아", team=[TeamMember(role="대표", name="김철수",
                                                         background="카카오 출신 PM 7년, 고려대학교 경영학 졸업")])
    leak = detail("bizplan", "- 대표: 카카오 출신 PM 7년, 고려대학교 경영학 졸업 (자사 자료)")
    exported = export_item(leak, "docx", profile)
    assert exported.notes == (blind_warning(leak.versions[-1].draft, profile),)
    assert "고려대학교" in exported.notes[0]
    assert export_item(detail("bizplan", "- 대표: ○○ 분야 7년 (자사 자료)"), "docx", profile).notes == ()
    assert export_item(detail("linkedin", "카카오 출신 PM 7년, 고려대학교 졸업"), "docx", profile).notes == ()
