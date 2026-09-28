"""Exporters (SPEC_V2 §6) against the real recorded sample-run drafts."""

from __future__ import annotations

import importlib.util
import io
import json
import os
import re
import struct
import subprocess
import sys
import zipfile
from html.parser import HTMLParser
from pathlib import Path
from xml.etree import ElementTree

import pytest

from insia_agents.channels import CHANNELS
from insia_agents.exporters import (ExportError, ExportFile, MissingDependencyError, capabilities, export_filename,
                                    export_item, export_run_zip, formats_for, slugify, sources_markdown)
from insia_agents.exporters.common import (ensure_hashtag_line, find_placeholders, inline_segments, item_date,
                                           parse_blocks, select_version, to_kst_date)
from insia_agents.exporters.instagram import Palette, parse_carousel, slides_html
from insia_agents.exporters.naver_html import extract_fragment, naver_fragment
from insia_agents.models import (Brief, ContentItem, ContentItemDetail, Draft, DraftVersion, Profile, ResearchPack,
                                 Review, TeamMember)
from insia_agents.storage import render_markdown

ROOT = Path(__file__).resolve().parents[1]
SAMPLE = ROOT / "examples" / "sample-run"
ROUNDS = {"bizplan": 1, "naver_blog": 1, "linkedin": 0, "instagram": 0}
RUN_ID = "20260927-231000-ab12"
HAS_DOCX = importlib.util.find_spec("docx") is not None
needs_docx = pytest.mark.skipif(not HAS_DOCX, reason='python-docx 미설치 (pip install "insia-smartagent[export]")')
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


@pytest.fixture(autouse=True)
def no_png_rendering(monkeypatch):
    """Keep the suite fast and deterministic; the PNG test opts back in."""
    monkeypatch.setenv("INSIA_RENDER", "0")
    monkeypatch.delenv("INSIA_CHROMIUM", raising=False)


def load_draft(channel: str) -> Draft:
    return Draft.model_validate_json((SAMPLE / "drafts" / f"{channel}.r{ROUNDS[channel]}.json").read_text("utf-8"))


def load_review(channel: str) -> Review | None:
    path = SAMPLE / "reviews" / f"{channel}.r{ROUNDS[channel]}.json"
    return Review.model_validate_json(path.read_text("utf-8")) if path.exists() else None


def sample_brief() -> Brief:
    return Brief.model_validate_json((SAMPLE / "brief.json").read_text("utf-8"))


def make_detail(channel: str, *, draft: Draft | None = None, run_id: str = RUN_ID, version: int = 2,
                created_at: str = "2026-09-27T16:30:00Z", scheduled_at: str = "", published_at: str = "",
                extra_versions: list[DraftVersion] | None = None, status: str = "draft") -> ContentItemDetail:
    draft = draft or load_draft(channel)
    item = ContentItem(id=f"it_{run_id}_{channel}", run_id=run_id, channel=channel, title=draft.title, version=version,
                       status=status, created_at=created_at, scheduled_at=scheduled_at, published_at=published_at)
    current = DraftVersion(id=f"dv_{channel}{version}", item_id=item.id, version=version, source="agent", draft=draft,
                           review=load_review(channel), created_at=created_at)
    return ContentItemDetail(item=item, versions=[*(extra_versions or []), current], brief=sample_brief())


def unzip(data: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


# ---------------------------------------------------------------------------
# Names, versions, dates
# ---------------------------------------------------------------------------


def test_slug_keeps_korean_and_strips_unsafe_chars():
    assert slugify("AI 마케팅 자동화, 1인 창업자가 도입 전 확인할 3가지") == "AI-마케팅-자동화-1인-창업자가-도입-전-확인할-3가지"
    assert slugify('a/b\\c:d*e?f"g<h>i|j_k') == "a-b-c-d-e-f-g-h-i-j-k"
    assert slugify("   ") == "초안"
    long = slugify("가" * 30 + " " + "나" * 30)
    assert len(long) <= 40 and not long.endswith("-")
    assert export_filename("2026-09-28", "linkedin", "훅: 첫 줄?", "txt") == "2026-09-28_linkedin_훅-첫-줄.txt"


def test_export_file_content_disposition_and_save(tmp_path):
    file = ExportFile("2026-09-28_naver_blog_AI-마케팅.html", "text/html; charset=utf-8", b"<p>x</p>")
    header = file.content_disposition()
    assert header.startswith('attachment; filename="2026-09-28_naver_blog_AI')
    assert "filename*=UTF-8''2026-09-28_naver_blog_AI-%EB%A7%88" in header
    header.encode("latin-1")  # must be a valid HTTP header value
    path = file.save(tmp_path / "exports")
    assert path.read_bytes() == b"<p>x</p>" and path.name == file.filename


def test_to_kst_date_and_item_date_priority():
    assert to_kst_date("2026-09-27T16:30:00Z") == "2026-09-28"  # 01:30 KST next day
    assert to_kst_date("2026-09-27T10:00:00+00:00") == "2026-09-27"
    assert to_kst_date("2026-10-02") == "2026-10-02"
    assert to_kst_date("nonsense") is None and to_kst_date("") is None
    assert item_date(make_detail("linkedin")) == "2026-09-28"
    assert item_date(make_detail("linkedin", scheduled_at="2026-10-05")) == "2026-10-05"
    assert item_date(make_detail("linkedin", scheduled_at="2026-10-05", published_at="2026-10-06T01:00:00Z")) == "2026-10-06"


def test_select_version_current_explicit_and_missing():
    old = load_draft("naver_blog").model_copy(update={"title": "예전 제목", "round": 0})
    detail = make_detail("naver_blog", version=2, extra_versions=[
        DraftVersion(id="dv_old", item_id=f"it_{RUN_ID}_naver_blog", version=1, source="agent", draft=old)])
    assert select_version(detail).version == 2
    assert select_version(detail, 1).draft.title == "예전 제목"
    assert "예전-제목" in export_item(detail, "md", version=1).filename
    with pytest.raises(ExportError, match="v7 버전을 찾을 수 없어요"):
        select_version(detail, 7)
    empty = detail.model_copy(update={"versions": []})
    with pytest.raises(ExportError, match="내보낼 초안이 아직 없어요"):
        export_item(empty, "md")


def test_format_matrix_and_korean_errors():
    assert formats_for("bizplan")[0] == "docx"
    assert formats_for("naver_blog")[0] == "html"
    assert formats_for("linkedin")[0] == "txt"
    assert formats_for("instagram")[0] == "zip"
    with pytest.raises(ExportError, match="링크드인 콘텐츠는 html 형식으로 내보낼 수 없어요. 가능한 형식: txt"):
        export_item(make_detail("linkedin"), "html")
    with pytest.raises(ExportError, match="지원하지 않는 형식이에요: pdf"):
        export_item(make_detail("linkedin"), "pdf")
    assert set(capabilities()) == {"docx", "png"}


def test_exporters_package_imports_optional_extras_lazily():
    code = ("import sys; sys.path.insert(0, 'src'); import insia_agents.exporters; "
            "print('docx' in sys.modules, 'playwright' in sys.modules)")
    out = subprocess.run([sys.executable, "-c", code], cwd=ROOT, capture_output=True, text=True, check=True)
    assert out.stdout.split() == ["False", "False"]


# ---------------------------------------------------------------------------
# Markdown parsing helpers
# ---------------------------------------------------------------------------


def test_parse_blocks_handles_gov_style_markers_tables_and_captions():
    text = ("# 제목\n\n## 1. 문제 인식\n\n□ (필요성) 첫 줄\n- 하위 항목 [s1]\n  - 더 깊은 항목\n  - ※ 주석\n"
            "※ 표 아래 주석\n\n< 1단계 집행계획 >\n\n| 비목 | 금액(원) |\n|---|---:|\n| 인건비 | 10,000,000 |\n"
            "| 합계 | 20,000,000 |\n\n[이미지: 흐름도]\n\n**굵게** 문단\n둘째 줄\n\n---\n")
    kinds = [(b.kind, b.marker, b.level) for b in parse_blocks(text)]
    assert kinds == [
        ("heading", "", 1), ("heading", "", 2), ("item", "□", 0), ("item", "-", 0), ("item", "-", 1),
        ("item", "-", 1), ("item", "※", 0), ("caption", "", 0), ("table", "", 0), ("image", "", 0),
        ("para", "", 0), ("rule", "", 0),
    ]
    table = parse_blocks(text)[8]
    assert table.header == ["비목", "금액(원)"] and table.rows[1] == ["합계", "20,000,000"]
    assert table.aligns == ["", "right"]


def test_placeholders_skip_citations_and_document_labels():
    line = "[대표자 성명] 외 [s12]·[s3], [○]명, [별첨 1] 양식, [확인 필요: 파일럿 업종] https://ex.com/a_(b)."
    kinds = [(s.kind, s.text) for s in inline_segments(line) if s.kind != "text"]
    assert kinds == [("placeholder", "[대표자 성명]"), ("placeholder", "[○]"),
                     ("placeholder", "[확인 필요: 파일럿 업종]"), ("url", "https://ex.com/a_(b)")]
    assert find_placeholders("[이미지: 사진] [문의 방법: 이메일]") == ["[문의 방법: 이메일]"]


def test_ensure_hashtag_line_never_drops_author_tags():
    assert ensure_hashtag_line("본문\n\n#a #b", ["#a", "#b"]) == "본문\n\n#a #b"
    assert ensure_hashtag_line("본문\n\n#a #x", ["#a", "#b"]) == "본문\n\n#a #x #b"
    assert ensure_hashtag_line("본문 끝.", ["#a", "#b"]) == "본문 끝.\n\n#a #b"
    assert ensure_hashtag_line("본문", []) == "본문"


# ---------------------------------------------------------------------------
# md / txt
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("channel", list(ROUNDS))
def test_markdown_matches_storage_render(channel):
    exported = export_item(make_detail(channel), "md")
    assert exported.content_type == "text/markdown; charset=utf-8"
    assert exported.filename.startswith(f"2026-09-28_{channel}_") and exported.filename.endswith(".md")
    assert exported.data.decode("utf-8") == render_markdown(load_draft(channel))


def test_linkedin_txt_is_exact_post_ending_with_hashtags():
    draft = load_draft("linkedin")
    text = export_item(make_detail("linkedin"), "txt").data.decode("utf-8")
    assert text == draft.content.strip() + "\n"  # sample already ends with its hashtag line
    assert text.rstrip().splitlines()[-1] == " ".join(draft.hashtags)
    no_tags = draft.model_copy(update={"content": draft.content.rsplit("\n\n", 1)[0]})
    text2 = export_item(make_detail("linkedin", draft=no_tags), "txt").data.decode("utf-8")
    assert text2.rstrip().splitlines()[-1] == " ".join(draft.hashtags)
    assert "http" not in text2


def test_instagram_txt_is_caption_only():
    draft = load_draft("instagram")
    text = export_item(make_detail("instagram"), "txt").data.decode("utf-8")
    assert "## 캐러셀" not in text and "슬라이드" not in text.split("\n")[0]
    assert text.startswith("혼자 콘텐츠를 다 쓰는 대표님께")
    assert text.rstrip().splitlines()[-1] == " ".join(draft.hashtags)
    assert len(text.strip()) <= CHANNELS["instagram"].limits["max_caption_chars"]
    broken = draft.model_copy(update={"content": "## 캐러셀\n\n### 슬라이드 1 — 훅\n- 문구: 안녕\n\n## 캡션\n\n"})
    with pytest.raises(ExportError, match="캡션"):
        export_item(make_detail("instagram", draft=broken), "txt")


def test_plain_txt_for_blog_and_bizplan_has_no_markdown_symbols():
    blog = export_item(make_detail("naver_blog"), "txt").data.decode("utf-8")
    assert blog.startswith("AI 마케팅 자동화, 1인 창업자가 도입 전 확인할 3가지\n")
    assert "## " not in blog and "**" not in blog
    assert blog.rstrip().splitlines()[-1] == " ".join(load_draft("naver_blog").hashtags)
    plan = export_item(make_detail("bizplan"), "txt").data.decode("utf-8")
    assert plan.startswith("INSIA 스마트에이전트 사업계획서\n") and "\n# " not in plan
    assert "인건비\t채용 예정 AI·백엔드 개발 인력 1명, 월 2,500,000원 × 4개월\t10,000,000" in plan


# ---------------------------------------------------------------------------
# Naver HTML
# ---------------------------------------------------------------------------


class _TagCollector(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.tags: list[str] = []
        self.attrs: list[str] = []

    def handle_starttag(self, tag, attrs):
        self.tags.append(tag)
        self.attrs.extend(name for name, _ in attrs)


def test_naver_fragment_is_smarteditor_friendly():
    draft = load_draft("naver_blog")
    fragment = naver_fragment(draft)
    parser = _TagCollector()
    parser.feed(fragment)
    assert set(parser.tags) <= {"p", "b", "br", "span", "a", "table", "tr", "td"}
    assert set(parser.attrs) <= {"style", "href", "target", "rel"}  # inline styles only, no classes
    assert fragment.count('<p style="font-size:1.25em;line-height:1.6;"><b>') == draft.content.count("\n## ")
    assert fragment.count("[이미지 자리]") == draft.content.count("[이미지:") == 4
    assert "border:2px dashed #03C75A" in fragment
    assert fragment.count("출처: ") == draft.content.count("출처: ")
    assert '<span style="background-color:#FFF3B0;">[대표 경험 추가:' in fragment
    assert fragment.rstrip().endswith(" ".join(draft.hashtags) + "</p>")
    assert "<b>1. 수치마다 출처와 기준 시점이 붙어 있나요?</b><br>" in fragment


def test_naver_html_export_is_standalone_preview_with_copy_buttons():
    exported = export_item(make_detail("naver_blog"), "html", Profile(banned_words=["최고"]))
    page = exported.data.decode("utf-8")
    assert exported.content_type == "text/html; charset=utf-8" and exported.filename.endswith(".html")
    assert page.startswith("<!doctype html>") and '<html lang="ko">' in page
    for button in ("copy-title", "copy-body", "copy-plain", "copy-tags"):
        assert f'id="{button}"' in page
    assert "ClipboardItem" in page and "execCommand('copy')" in page
    assert "게시 전에 채우거나 지워야 할 자리 3곳" in page
    assert "형식 점검 (7/7 통과)" in page  # 6 channel checks + banned-word check from the profile
    assert extract_fragment(page) == naver_fragment(load_draft("naver_blog"))
    data = json.loads(re.search(r'<script type="application/json" id="insia-data">(.*?)</script>', page, re.S).group(1))
    assert data["title"] == load_draft("naver_blog").title


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def _check_child_order(xml: bytes) -> None:
    from insia_agents.exporters.docx_writer import ORDERS

    root = ElementTree.fromstring(xml)
    for parent in root.iter():
        local = parent.tag.split("}")[-1]
        if local not in ORDERS or not parent.tag.startswith(W):
            continue
        if local == "rPr" and parent.find(f"{W}ins") is not None:
            continue
        order = ORDERS[local]
        names = [c.tag.split("}")[-1] for c in parent if c.tag.startswith(W)]
        known = [order.index(n) for n in names if n in order]
        assert known == sorted(known), f"<w:{local}> children out of schema order: {names}"


def _docx_parts(data: bytes) -> dict[str, bytes]:
    parts = unzip(data)
    assert "word/document.xml" in parts and "word/styles.xml" in parts
    return parts


@needs_docx
def test_bizplan_docx_structure_for_word_and_hangul():
    import docx
    from docx.shared import Cm

    exported = export_item(make_detail("bizplan"), "docx")
    assert exported.content_type.endswith("wordprocessingml.document") and exported.filename.endswith(".docx")
    document = docx.Document(io.BytesIO(exported.data))
    section = document.sections[0]
    assert abs(section.page_width - Cm(21.0)) < Cm(0.01) and abs(section.page_height - Cm(29.7)) < Cm(0.01)
    for margin in (section.left_margin, section.right_margin, section.top_margin, section.bottom_margin):
        assert abs(margin - Cm(2.0)) < Cm(0.01)

    paragraphs = document.paragraphs
    assert paragraphs[0].style.name == "Title" and paragraphs[0].text == "INSIA 스마트에이전트 사업계획서"
    assert "검수" in paragraphs[1].text and "v2" in paragraphs[1].text
    h1 = [p.text for p in paragraphs if p.style.name == "Heading 1"]
    assert any("문제 인식" in t for t in h1) and any("팀 구성" in t for t in h1) and "참고자료" in h1
    assert any(p.style.name == "Heading 2" and p.text.startswith("2-1.") for p in paragraphs)

    square = next(p for p in paragraphs if p.text.startswith("□ (필요성)"))
    assert square.paragraph_format.first_line_indent < 0 < square.paragraph_format.left_indent
    assert square.runs[0].bold and square.runs[1].text == "(필요성)" and square.runs[1].bold
    dash = next(p for p in paragraphs if p.text.startswith("- 소상공인 기업체 613.4만개"))
    nested = next(p for p in paragraphs if p.text.startswith("- 종사자 961.0만명"))
    assert square.paragraph_format.left_indent < dash.paragraph_format.left_indent < nested.paragraph_format.left_indent
    note = next(p for p in paragraphs if p.text.startswith("※ 가정: 단가는"))
    assert note.runs[-1].font.size.pt == 9.5

    md_tables = sum(1 for b in parse_blocks(load_draft("bizplan").content) if b.kind == "table")
    assert len(document.tables) == md_tables >= 10
    budget = next(t for t in document.tables if t.rows[0].cells[0].text == "비목")
    assert budget.rows[0].cells[2].text == "정부지원사업비(원)"
    assert budget.rows[-1].cells[2].text == "20,000,000"
    assert budget.rows[1].cells[2].paragraphs[0].alignment == 2  # right-aligned amounts

    parts = _docx_parts(exported.data)
    body = parts["word/document.xml"].decode("utf-8")
    assert body.count("<w:tblHeader/>") == md_tables
    assert 'w:fill="DCE3F0"' in body and "<w:tblBorders>" in body and 'w:type="fixed"' in body
    assert '<w:highlight w:val="yellow"/>' in body  # [대표자 성명] etc.
    assert "<w:hyperlink" in body and "https://www.mss.go.kr" in parts["word/_rels/document.xml.rels"].decode()
    styles = parts["word/styles.xml"].decode("utf-8")
    assert 'w:eastAsia="맑은 고딕"' in styles and 'w:ascii="Malgun Gothic"' in styles
    assert 'w:eastAsia="ko-KR"' in styles and "eastAsiaTheme" not in styles.split("</w:docDefaults>")[0]
    footer = next(v.decode("utf-8") for k, v in parts.items() if k.startswith("word/footer"))
    assert "INSIA 초안 — 제출 전 사람 검토 필수" in footer and "PAGE" in footer and "NUMPAGES" in footer
    for name in ("word/document.xml", "word/styles.xml"):
        _check_child_order(parts[name])
    # the blind rule is not triggered without a profile
    assert "블라인드" not in body and exported.notes == ()


@needs_docx
def test_bizplan_docx_prints_blind_check_when_team_names_are_exposed():
    import docx

    draft = load_draft("bizplan")
    exposed = draft.model_copy(update={"content": draft.content.replace("[대표자 성명], [학위·전공]", "홍길동, [학위·전공]")})
    profile = Profile(team=[TeamMember(role="대표", name="홍길동"), TeamMember(role="CTO", name="김철수")])
    exported = export_item(make_detail("bizplan", draft=exposed), "docx", profile)
    text = "\n".join(c.text for t in docx.Document(io.BytesIO(exported.data)).tables for row in t.rows for c in row.cells)
    assert "블라인드 확인 필요 — 팀원 실명 1개가 본문에 보여요: 홍길동" in text
    assert exported.notes and "홍길동" in exported.notes[0] and "김철수" not in exported.notes[0]
    clean = export_item(make_detail("bizplan"), "docx", profile)
    assert "블라인드" not in unzip(clean.data)["word/document.xml"].decode("utf-8") and clean.notes == ()


@needs_docx
@pytest.mark.parametrize("channel", ["naver_blog", "linkedin", "instagram"])
def test_docx_works_for_other_channels_as_simple_document(channel):
    import docx

    draft = load_draft(channel)
    exported = export_item(make_detail(channel), "docx")
    document = docx.Document(io.BytesIO(exported.data))
    assert document.paragraphs[0].style.name == "Title" and document.paragraphs[0].text == draft.title
    text = "\n".join(p.text for p in document.paragraphs)
    assert " ".join(draft.hashtags) in text
    parts = _docx_parts(exported.data)
    for name in ("word/document.xml", "word/styles.xml"):
        _check_child_order(parts[name])
    if channel == "naver_blog":
        boxes = [t for t in document.tables if t.rows[0].cells[0].text.startswith("[이미지 자리]")]
        assert len(boxes) == 4
        assert sum(1 for p in document.paragraphs if p.style.name == "Heading 1") == draft.content.count("\n## ")
    if channel == "instagram":
        assert sum(1 for p in document.paragraphs if p.style.name == "Heading 2") == 9
        label = next(p for p in document.paragraphs if p.text.startswith("- 문구:"))
        assert label.runs[1].text == "문구:" and label.runs[1].bold


def test_docx_without_python_docx_raises_install_hint(monkeypatch):
    monkeypatch.setitem(sys.modules, "docx", None)
    with pytest.raises(MissingDependencyError) as info:
        export_item(make_detail("bizplan"), "docx")
    assert 'pip install "insia-smartagent[export]"' in str(info.value)
    assert info.value.extra == "export" and info.value.package == "python-docx"
    # the per-item zip still bundles what it can and explains what is missing
    bundle = export_item(make_detail("bizplan"), "zip")
    names = unzip(bundle.data)
    assert any(n.endswith(".md") for n in names) and not any(n.endswith(".docx") for n in names)
    assert any("python-docx" in n for n in bundle.notes)


# ---------------------------------------------------------------------------
# Instagram
# ---------------------------------------------------------------------------


def test_parse_carousel_reads_all_slide_fields():
    slides = parse_carousel(load_draft("instagram").content)
    assert len(slides) == 9 and [s.number for s in slides] == list(range(1, 10))
    first = slides[0]
    assert first.title == "훅: 혼자 다 쓰는 대표님께"
    assert first.parts == ["블로그·인스타·사업계획서, 혼자 다 쓰고 있다면", "AI와 나눠 맡는 4단계 순서, 끝까지 넘겨 보세요"]
    assert first.visual.startswith("(전 장 공통 체계)") and first.alt.startswith("인디고 배경에")
    assert slides[1].source.startswith("중소벤처기업부·소상공인시장진흥공단") and not slides[0].source
    assert len(slides[8].parts) == 3


def test_palette_uses_brand_colors_with_readable_text():
    assert Palette.from_profile(None) == Palette("#0B1220", "#6D5EF5", "#F5C451")
    assert Palette.from_profile(Profile(brand_colors=["nope", "#12"])) == Palette("#0B1220", "#6D5EF5", "#F5C451")
    brand = Palette.from_profile(Profile(brand_colors=["#03c75a", "#111111"]))
    assert brand.primary == "#03C75A" and brand.accent == "#111111" and brand.dark == "#111111"
    light = Palette.from_profile(Profile(brand_colors=["#FFE812"]))
    assert light.theme("cover")["--fg"] == "#111827"  # dark text on a light brand color
    assert Palette.from_profile(None).theme("inner")["--fg"] == "#FFFFFF"


def test_slides_html_has_counter_brand_handle_and_sources():
    slides = parse_carousel(load_draft("instagram").content)
    page = slides_html(slides, Profile(service_name="INSIA 스마트에이전트", instagram_handle="https://instagram.com/insia.ai/",
                                        brand_colors=["#0A66C2"]), title="t")
    assert page.count('<section class="slide') == 9
    assert "width: 1080px; height: 1350px" in page and "@page { size: 1080px 1350px" in page
    assert '<span class="counter">1/9</span>' in page and '<span class="counter">9/9</span>' in page
    assert "@insia.ai" in page and "INSIA 스마트에이전트</span>" in page and "#0A66C2" in page
    assert "평균 <em>1.57명</em>" in page
    assert page.count('<p class="source">출처: ') == sum(1 for s in slides if s.source)
    assert '<p class="pill">INSIA 스마트에이전트는 프로필 링크에서</p>' in page
    assert "word-break: keep-all" in page and "Malgun Gothic" in page


def test_instagram_zip_falls_back_to_slides_html_with_readme():
    draft = load_draft("instagram")
    exported = export_item(make_detail("instagram"), "zip", Profile(instagram_handle="insia.ai"))
    assert exported.content_type == "application/zip" and exported.filename.endswith(".zip")
    files = unzip(exported.data)
    assert set(files) == {"slides.html", "caption.txt", "alt-text.txt", "README.txt"}
    assert files["caption.txt"].decode("utf-8") == export_item(make_detail("instagram"), "txt").data.decode("utf-8")
    alt = files["alt-text.txt"].decode("utf-8")
    assert alt.count("[slide-") == 9 and "[slide-09.png] 슬라이드 9" in alt
    readme = files["README.txt"].decode("utf-8")
    assert "slides.html을 대신 넣었어요" in readme and "INSIA_RENDER=0" in readme
    assert "디자인 메모" in readme and draft.title in readme
    assert exported.notes and "slides.html" in exported.notes[0]
    assert files["slides.html"].decode("utf-8").count('<section class="slide') == 9


def test_instagram_zip_without_slides_is_a_clear_error():
    draft = load_draft("instagram").model_copy(update={"content": "## 캡션\n첫 줄\n#a #b #c"})
    with pytest.raises(ExportError, match="캐러셀 슬라이드"):
        export_item(make_detail("instagram", draft=draft), "zip")


def _chromium() -> str | None:
    for candidate in (os.environ.get("INSIA_CHROMIUM"), "/opt/pw-browsers/chromium"):
        if candidate and Path(candidate).exists():
            return candidate
    return None


CHROMIUM = _chromium()


@pytest.mark.skipif(importlib.util.find_spec("playwright") is None or CHROMIUM is None,
                    reason="Playwright/Chromium 없음 — PNG 렌더링 테스트 생략")
def test_instagram_zip_renders_1080x1350_pngs(monkeypatch):
    monkeypatch.delenv("INSIA_RENDER", raising=False)
    monkeypatch.setenv("INSIA_CHROMIUM", CHROMIUM)
    exported = export_item(make_detail("instagram"), "zip", Profile(instagram_handle="insia.ai"))
    files = unzip(exported.data)
    pngs = sorted(n for n in files if n.endswith(".png"))
    assert pngs == [f"slide-{i:02d}.png" for i in range(1, 10)]
    for name in pngs:
        head = files[name][:24]
        assert head[:8] == b"\x89PNG\r\n\x1a\n"
        assert struct.unpack(">II", head[16:24]) == (1080, 1350)
    assert "slides.html" not in files and exported.notes == ()
    assert "slide-01.png ~ slide-09.png" in files["README.txt"].decode("utf-8")


@pytest.mark.skipif(importlib.util.find_spec("playwright") is None or CHROMIUM is None
                    or not Path("/usr/share/fonts/truetype/dejavu").is_dir(),
                    reason="Playwright/Chromium 또는 DejaVu 글꼴 없음")
def test_missing_hangul_font_falls_back_instead_of_rendering_boxes(monkeypatch, tmp_path):
    conf = tmp_path / "fonts.conf"
    conf.write_text('<?xml version="1.0"?><fontconfig><dir>/usr/share/fonts/truetype/dejavu</dir>'
                    f"<cachedir>{tmp_path / 'cache'}</cachedir></fontconfig>", encoding="utf-8")
    monkeypatch.setenv("FONTCONFIG_FILE", str(conf))  # Chromium inherits the environment
    monkeypatch.delenv("INSIA_RENDER", raising=False)
    monkeypatch.setenv("INSIA_CHROMIUM", CHROMIUM)
    exported = export_item(make_detail("instagram"), "zip")
    assert "slides.html" in unzip(exported.data)
    assert "한글 글꼴" in exported.notes[0]


needs_playwright = pytest.mark.skipif(importlib.util.find_spec("playwright") is None, reason="Playwright 없음")


@needs_playwright
def test_hung_browser_hits_deadline_and_falls_back(monkeypatch):
    import time

    from insia_agents.exporters import instagram

    monkeypatch.delenv("INSIA_RENDER", raising=False)
    monkeypatch.setattr(instagram, "RENDER_DEADLINE_S", 0.2)
    monkeypatch.setattr(instagram, "_render_sync", lambda *a: time.sleep(1.5))
    exported = export_item(make_detail("instagram"), "zip")
    assert "slides.html" in unzip(exported.data) and "초가 넘게 걸려" in exported.notes[0]


@needs_playwright
def test_render_works_from_inside_an_asyncio_loop(monkeypatch):
    import asyncio

    from insia_agents.exporters import instagram

    monkeypatch.delenv("INSIA_RENDER", raising=False)
    fake_png = b"\x89PNG\r\n\x1a\n" + struct.pack(">I", 13) + b"IHDR" + struct.pack(">II", 1080, 1350)
    monkeypatch.setattr(instagram, "_render_sync", lambda html, count, exe: [fake_png] * count)

    async def run():
        return export_item(make_detail("instagram"), "zip")

    files = unzip(asyncio.run(run()).data)
    assert sorted(n for n in files if n.endswith(".png"))[-1] == "slide-09.png"


def test_bad_chromium_path_falls_back_with_reason(monkeypatch):
    monkeypatch.delenv("INSIA_RENDER", raising=False)
    monkeypatch.setenv("INSIA_CHROMIUM", "/nonexistent/chrome")
    exported = export_item(make_detail("instagram"), "zip")
    assert "slides.html" in unzip(exported.data)
    if importlib.util.find_spec("playwright") is not None:
        assert "INSIA_CHROMIUM" in exported.notes[0]


# ---------------------------------------------------------------------------
# Per-item zip for other channels, run zip
# ---------------------------------------------------------------------------


def test_linkedin_zip_bundles_every_format():
    names = set(unzip(export_item(make_detail("linkedin"), "zip").data))
    exts = {n.rsplit(".", 1)[1] for n in names}
    assert {"txt", "md"} <= exts and (("docx" in exts) == HAS_DOCX)


class FakeWorkspace:
    """Only the Workspace methods export_run_zip may use (SPEC_V2 §2)."""

    def __init__(self, run: dict, details: list[ContentItemDetail], profile: Profile | None = None) -> None:
        self.run, self.details, self.profile = run, details, profile
        self.calls: list[str] = []

    def get_run(self, run_id: str) -> dict | None:
        self.calls.append("get_run")
        return self.run if run_id == self.run["run_id"] else None

    def list_items(self, *, status=None, channel=None, limit=200):
        self.calls.append("list_items")
        return [d.item for d in self.details][:limit]

    def get_item(self, item_id: str):
        self.calls.append("get_item")
        return next((d for d in self.details if d.item.id == item_id), None)

    def get_profile(self) -> Profile:
        self.calls.append("get_profile")
        return self.profile or Profile()


def _run_dict() -> dict:
    research = ResearchPack.model_validate_json((SAMPLE / "research.json").read_text("utf-8"))
    return {"run_id": RUN_ID, "kind": "pipeline", "status": "completed", "started_at": "2026-09-27T23:10:00Z",
            "brief": sample_brief().model_dump(mode="json"), "research": research.model_dump(mode="json"),
            "profile": Profile(instagram_handle="snapshot.handle").model_dump(mode="json"), "cost_usd": 1.2345}


def test_run_zip_contains_all_channel_exports_and_research():
    other = make_detail("linkedin", run_id="20260101-000000-zzzz")
    details = [make_detail(ch) for ch in ROUNDS] + [other]
    workspace = FakeWorkspace(_run_dict(), details, Profile(service_name="INSIA 스마트에이전트"))
    exported = export_run_zip(workspace, RUN_ID)
    root = f"2026-09-28_run_{RUN_ID}"
    assert exported.filename == f"{root}.zip" and exported.content_type == "application/zip"
    files = unzip(exported.data)
    assert all(name.startswith(root + "/") for name in files)
    names = {n[len(root) + 1:] for n in files}
    assert {"README.txt", "research.json", "sources.md", "brief.json"} <= names
    assert any(n.startswith("naver_blog/") and n.endswith(".html") for n in names)
    assert any(n.startswith("linkedin/") and n.endswith(".txt") for n in names)
    assert {"instagram/caption.txt", "instagram/alt-text.txt", "instagram/slides.html", "instagram/README.txt"} <= names
    assert any(n.startswith("bizplan/") and n.endswith(".docx") for n in names) == HAS_DOCX
    assert sum(1 for n in names if n.endswith(".md") and "/" in n) == 4
    assert {f"{ch}/review.json" for ch in ROUNDS} <= names
    assert not any("20260101" in n for n in names)  # items of other runs stay out
    research = json.loads(files[f"{root}/research.json"])
    assert len(research["sources"]) == len(ResearchPack.model_validate(research).sources) > 0
    readme = files[f"{root}/README.txt"].decode("utf-8")
    assert f"실행 ID: {RUN_ID}" in readme and "API 비용: 약 $1.23" in readme
    assert "- 사업계획서: \"INSIA 스마트에이전트 사업계획서\" — v2" in readme and "사람이 검토하고 승인" in readme
    sources = files[f"{root}/sources.md"].decode("utf-8")
    assert sources.startswith("# 출처 목록 — 1인 창업자") and "## Tier 1 · 공식·공공 자료" in sources
    assert "[s1] " in sources and "## 근거 (findings)" in sources
    assert set(workspace.calls) <= {"get_run", "list_items", "get_item", "get_profile"}


def test_run_zip_uses_run_profile_snapshot_and_reports_unknown_runs():
    class NoProfileWorkspace(FakeWorkspace):
        get_profile = None  # type: ignore[assignment]

    workspace = NoProfileWorkspace(_run_dict(), [make_detail("instagram")])
    files = unzip(export_run_zip(workspace, RUN_ID).data)
    slides = next(v for k, v in files.items() if k.endswith("instagram/slides.html")).decode("utf-8")
    assert "@snapshot.handle" in slides
    with pytest.raises(ExportError, match="실행 기록을 찾을 수 없어요"):
        export_run_zip(workspace, "missing-run")


def test_sources_markdown_groups_user_materials_and_handles_missing_research():
    pack = ResearchPack.model_validate({
        "findings": [{"id": "f1", "question_id": "q1", "claim": "베타 사용자 120명", "source_ids": ["s2"],
                      "confidence": "low"}],
        "sources": [{"id": "s1", "title": "통계", "url": "https://kosis.kr", "tier": 1, "publisher": "통계청"},
                    {"id": "s2", "title": "회사 소개서", "url": "user://u3", "tier": 1, "origin": "user",
                     "publisher": "사용자 제공 자료"}],
        "gaps": ["시장 규모 최신치"],
    })
    text = sources_markdown(pack, topic="테스트")
    assert "## 사용자 제공 자료 (1)" in text and "사용자 제공 자료 (user://u3)" in text
    assert "## Tier 1 · 공식·공공 자료 (1)" in text and "- [f1] 베타 사용자 120명 — [s2], 신뢰도 낮음" in text
    assert "## 확인하지 못한 부분" in text
    assert "리서치 기록이 없어요" in sources_markdown(None)


def test_run_zip_with_real_workspace(tmp_path):
    db = pytest.importorskip("insia_agents.db")
    from insia_agents.models import ChannelResult

    workspace = db.Workspace(tmp_path / "ws")
    brief = sample_brief()
    workspace.create_run(RUN_ID, brief, mode="mock", model="claude-opus-5", profile=Profile(instagram_handle="insia.ai"))
    workspace.update_run(RUN_ID, research=ResearchPack.model_validate_json((SAMPLE / "research.json").read_text("utf-8")),
                         status="completed")
    for channel, rounds in ROUNDS.items():
        drafts = [Draft.model_validate_json((SAMPLE / "drafts" / f"{channel}.r{i}.json").read_text("utf-8"))
                  for i in range(rounds + 1)]
        reviews = [Review.model_validate_json(p.read_text("utf-8"))
                   for p in (SAMPLE / "reviews" / f"{channel}.r{i}.json" for i in range(rounds + 1)) if p.exists()]
        workspace.upsert_item_from_result(RUN_ID, ChannelResult(channel=channel, final=drafts[-1], drafts=drafts,
                                                                reviews=reviews, passed=True, rounds=rounds), brief)
    exported = export_run_zip(workspace, RUN_ID)
    names = {n.split("/", 1)[1] for n in unzip(exported.data)}
    assert {"research.json", "sources.md", "README.txt", "instagram/caption.txt"} <= names
    assert {f"{ch}/review.json" for ch in ROUNDS} <= names
    assert any(n.startswith("naver_blog/") and n.endswith(".html") for n in names)
    detail = workspace.get_item(f"it_{RUN_ID}_linkedin")
    assert export_item(detail, "txt").data.decode("utf-8").rstrip().endswith("#소상공인")
