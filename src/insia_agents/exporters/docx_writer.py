"""Word (.docx) export built with python-docx (optional extra ``[export]``).

Targets both MS Word and Hancom Hangul (한글):
- A4, 2.0 cm margins, Korean font (맑은 고딕 / Malgun Gothic) set on the
  document defaults and every style with explicit ``w:eastAsia`` rFonts so
  theme fonts never fall back to a Japanese/Latin face.
- Headings as real Word heading styles (navigation pane / 개요 보기).
- 개조식 lines (□ / - / ※ / →) as separate paragraphs with hanging indents so
  wrapped lines align under the text, markers kept as written.
- Markdown tables → real Word tables: explicit borders, shaded header row that
  repeats on page breaks, fixed column widths (Hangul ignores autofit).
- Fill-in placeholders ([대표자 성명], [확인 필요: …], [○]) highlighted yellow.
- Footer: "INSIA 초안 — 제출 전 사람 검토 필수 | N / M".

python-docx is imported lazily; a missing install raises
``MissingDependencyError`` with the pip command.
"""

from __future__ import annotations

import importlib
import io
import re
import unicodedata
from datetime import datetime, timezone
from urllib.parse import quote

from ..models import Draft, Profile
from .common import (EXTRA_HINT, TOP_LEVEL_MARKERS, Block, MissingDependencyError, inline_segments, is_hashtag_line,
                     parse_blocks, strip_inline)

FONT_LATIN = "Malgun Gothic"
FONT_EAST_ASIA = "맑은 고딕"
PAGE_WIDTH_CM = 21.0
PAGE_HEIGHT_CM = 29.7
MARGIN_CM = 2.0
TEXT_WIDTH_CM = PAGE_WIDTH_CM - 2 * MARGIN_CM
FOOTER_TEXT = "INSIA 초안 — 제출 전 사람 검토 필수"

INK = "1F2328"
MUTED = "5B616B"
NOTE = "4B5563"
HEADER_FILL = "DCE3F0"
BORDER = "8A94A6"

# Canonical child order (ECMA-376) for the containers we touch. Word refuses
# ("unreadable content") files whose children are out of order.
PPR_ORDER = ("pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
             "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
             "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid",
             "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
             "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange")
RPR_ORDER = ("rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline",
             "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing", "w",
             "kern", "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign",
             "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath")
TBLPR_ORDER = ("tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
               "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook",
               "tblCaption", "tblDescription", "tblPrChange")
TCPR_ORDER = ("cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
              "textDirection", "tcFitText", "vAlign", "hideMark", "headers", "cellIns", "cellDel", "cellMerge",
              "tcPrChange")
TRPR_ORDER = ("cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter", "cantSplit", "trHeight",
              "tblHeader", "tblCellSpacing", "jc", "hidden", "ins", "del", "trPrChange")
ORDERS = {"pPr": PPR_ORDER, "rPr": RPR_ORDER, "tblPr": TBLPR_ORDER, "tcPr": TCPR_ORDER, "trPr": TRPR_ORDER}

_NUMERIC_CELL = re.compile(r"^(약\s?)?[-+]?[\d,.]+\s?(원|%|%p|개|명|건|회|억원|만원|백만원|천원|달러)?$")
_BR = re.compile(r"<br\s*/?>", re.IGNORECASE)
_LABEL_PAREN = re.compile(r"^(\([^()\n]{1,16}\))(\s*)")
_LABEL_COLON = re.compile(r"^([^\s:：\[\]]{1,10})([:：])(\s)")
_REFERENCE_HEADINGS = ("참고자료", "참고 자료", "참고문헌", "출처")


def require_docx():
    """Import python-docx or raise a Korean install hint."""
    try:
        return importlib.import_module("docx")
    except ImportError as exc:
        raise MissingDependencyError(
            "Word(.docx)로 내보내려면 python-docx가 필요해요. "
            f"`{EXTRA_HINT.format(extra='export')}`로 설치한 뒤 다시 시도해 주세요.",
            package="python-docx", extra="export",
        ) from exc


def display_width(text: str) -> int:
    return sum(2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1 for ch in text)


def column_widths(rows: list[list[str]], total_cm: float = TEXT_WIDTH_CM, min_cm: float = 1.4) -> list[float]:
    """Fixed column widths (cm) from content length; wide (Korean) chars count double."""
    ncols = max(len(r) for r in rows)
    weights: list[float] = []
    for c in range(ncols):
        lens = [display_width(strip_inline(_BR.sub(" ", r[c]))) if c < len(r) else 0 for r in rows]
        longest = max(lens) if lens else 0
        average = sum(lens) / len(lens) if lens else 0
        weights.append(max(4, min(longest, 60)) * 0.5 + max(4, min(average, 60)) * 0.5)
    raw = [max(min_cm, total_cm * w / sum(weights)) for w in weights]
    scale = total_cm / sum(raw)
    return [round(w * scale, 2) for w in raw]


def _numeric_columns(rows: list[list[str]]) -> set[int]:
    numeric: set[int] = set()
    if not rows:
        return numeric
    for c in range(max(len(r) for r in rows)):
        values = [strip_inline(r[c]).strip() for r in rows if c < len(r) and strip_inline(r[c]).strip()]
        if values and sum(1 for v in values if _NUMERIC_CELL.match(v)) / len(values) >= 0.6:
            numeric.add(c)
    return numeric


class _Builder:
    def __init__(self, draft: Draft, *, meta: str, profile: Profile | None, channel_label: str) -> None:
        docx = require_docx()
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
        from docx.opc.constants import RELATIONSHIP_TYPE
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt, RGBColor

        self.Cm, self.Pt, self.RGBColor = Cm, Pt, RGBColor
        self.OxmlElement, self.qn = OxmlElement, qn
        self.ALIGN, self.HIGHLIGHT = WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
        self.TABLE_ALIGN, self.VALIGN = WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
        self.RT = RELATIONSHIP_TYPE
        self.doc = docx.Document()
        self.draft = draft
        self.meta = meta
        self.profile = profile
        self.channel_label = channel_label
        self.in_references = False
        self.references_level = 0
        self.last_was_table = False

    # -- low-level XML helpers ----------------------------------------------

    def _el(self, tag: str, **attrs: str):
        el = self.OxmlElement(tag)
        for key, value in attrs.items():
            el.set(self.qn(f"w:{key}"), value)
        return el

    def _insert(self, parent, child) -> None:
        """Insert ``child`` into ``parent`` respecting the schema order (replacing a same-tag child)."""
        order = ORDERS[parent.tag.split("}")[1]]
        name = child.tag.split("}")[1]
        for existing in parent.findall(child.tag):
            parent.remove(existing)
        index = order.index(name)
        for position, sibling in enumerate(parent):
            sibling_name = sibling.tag.split("}")[1] if isinstance(sibling.tag, str) else ""
            if sibling_name in order and order.index(sibling_name) > index:
                parent.insert(position, child)
                return
        parent.append(child)

    def _set_fonts(self, rpr) -> None:
        rfonts = rpr.get_or_add_rFonts()
        for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            rfonts.attrib.pop(self.qn(f"w:{attr}"), None)
        rfonts.set(self.qn("w:ascii"), FONT_LATIN)
        rfonts.set(self.qn("w:hAnsi"), FONT_LATIN)
        rfonts.set(self.qn("w:eastAsia"), FONT_EAST_ASIA)
        rfonts.set(self.qn("w:cs"), FONT_LATIN)

    def _set_color(self, font, hex_color: str) -> None:
        font.color.rgb = self.RGBColor.from_string(hex_color)
        color = font.element.rPr.find(self.qn("w:color")) if font.element.rPr is not None else None
        if color is not None:
            for attr in ("themeColor", "themeShade", "themeTint"):
                color.attrib.pop(self.qn(f"w:{attr}"), None)

    # -- document setup ------------------------------------------------------

    def setup(self) -> None:
        Cm, Pt = self.Cm, self.Pt
        section = self.doc.sections[0]
        section.page_width, section.page_height = Cm(PAGE_WIDTH_CM), Cm(PAGE_HEIGHT_CM)
        section.left_margin = section.right_margin = Cm(MARGIN_CM)
        section.top_margin = section.bottom_margin = Cm(MARGIN_CM)
        section.header_distance = section.footer_distance = Cm(1.0)

        styles = self.doc.styles
        defaults = styles.element.find(self.qn("w:docDefaults"))
        if defaults is not None:
            rpr = defaults.find(f"{self.qn('w:rPrDefault')}/{self.qn('w:rPr')}")
            if rpr is not None:
                self._set_fonts(rpr)
                lang = rpr.find(self.qn("w:lang"))
                if lang is None:
                    lang = self._el("w:lang")
                    self._insert(rpr, lang)
                lang.set(self.qn("w:val"), "en-US")
                lang.set(self.qn("w:eastAsia"), "ko-KR")

        normal = styles["Normal"]
        normal.font.size = Pt(10.5)
        self._set_fonts(normal.element.get_or_add_rPr())
        self._set_color(normal.font, INK)
        fmt = normal.paragraph_format
        fmt.space_before, fmt.space_after, fmt.line_spacing = Pt(0), Pt(3), 1.3

        heading_specs = {
            "Title": (20, "111827", 0, 4),
            "Heading 1": (15, "1B2A4A", 18, 6),
            "Heading 2": (12.5, "243B6B", 12, 4),
            "Heading 3": (11, "333333", 9, 3),
        }
        for name, (size, color, before, after) in heading_specs.items():
            style = styles[name]
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.italic = False
            self._set_fonts(style.element.get_or_add_rPr())
            self._set_color(style.font, color)
            pf = style.paragraph_format
            pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), 1.25
            pf.keep_with_next = True
            ppr = style.element.get_or_add_pPr()
            for border in ppr.findall(self.qn("w:pBdr")):
                ppr.remove(border)
        styles["Title"].paragraph_format.alignment = self.ALIGN.CENTER
        h1_ppr = styles["Heading 1"].element.get_or_add_pPr()
        border = self._el("w:pBdr")
        border.append(self._el("w:bottom", val="single", sz="6", space="4", color="C9D1E0"))
        self._insert(h1_ppr, border)

        props = self.doc.core_properties
        props.title = self.draft.title
        props.subject = f"{self.channel_label} 초안"
        props.author = "INSIA 스마트에이전트"
        props.last_modified_by = "INSIA 스마트에이전트"
        props.comments = "AI 초안 — 제출·게시 전 사람 검토 필수"
        props.language = "ko-KR"
        now = datetime.now(timezone.utc).replace(microsecond=0, tzinfo=None)
        props.created = props.modified = now
        props.revision = 1

        self._footer(section)

    def _field(self, paragraph, instruction: str, size: float) -> None:
        for kind in ("begin", "instr", "separate", "text", "end"):
            run = paragraph.add_run()
            run.font.size = self.Pt(size)
            self._set_color(run.font, MUTED)
            if kind == "instr":
                instr = self.OxmlElement("w:instrText")
                instr.set(self.qn("xml:space"), "preserve")
                instr.text = f" {instruction} "
                run._r.append(instr)
            elif kind == "text":
                run.text = "1"
            else:
                run._r.append(self._el("w:fldChar", fldCharType=kind))

    def _footer(self, section) -> None:
        footer = section.footer
        footer.is_linked_to_previous = False
        paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        paragraph.style = self.doc.styles["Normal"]
        paragraph.alignment = self.ALIGN.CENTER
        paragraph.paragraph_format.space_after = self.Pt(0)
        run = paragraph.add_run(f"{FOOTER_TEXT}   |   ")
        run.font.size = self.Pt(8.5)
        self._set_color(run.font, MUTED)
        self._field(paragraph, "PAGE", 8.5)
        sep = paragraph.add_run(" / ")
        sep.font.size = self.Pt(8.5)
        self._set_color(sep.font, MUTED)
        self._field(paragraph, "NUMPAGES", 8.5)

    # -- runs ------------------------------------------------------------------

    def _hyperlink(self, paragraph, url: str, size: float | None) -> None:
        target = quote(url, safe=":/?#[]@!$&'()*+,;=%~-._")
        r_id = paragraph.part.relate_to(target, self.RT.HYPERLINK, is_external=True)
        link = self.OxmlElement("w:hyperlink")
        link.set(self.qn("r:id"), r_id)
        link.set(self.qn("w:history"), "1")
        run = self.OxmlElement("w:r")
        rpr = self.OxmlElement("w:rPr")
        rpr.append(self._el("w:color", val="0563C1"))
        if size:
            half_points = str(int(round(size * 2)))
            rpr.append(self._el("w:sz", val=half_points))
            rpr.append(self._el("w:szCs", val=half_points))
        rpr.append(self._el("w:u", val="single"))
        run.append(rpr)
        text = self.OxmlElement("w:t")
        text.set(self.qn("xml:space"), "preserve")
        text.text = url
        run.append(text)
        link.append(run)
        paragraph._p.append(link)

    def _runs(self, paragraph, text: str, *, size: float | None = None, color: str | None = None,
              bold: bool = False, italic: bool = False, links: bool = True) -> None:
        pieces = _BR.split(text)
        for index, piece in enumerate(pieces):
            if index:
                paragraph.add_run().add_break()
            for seg in inline_segments(piece):
                if seg.kind == "url" and links:
                    self._hyperlink(paragraph, seg.text, size)
                    continue
                run = paragraph.add_run(seg.text)
                if bold or seg.bold:
                    run.bold = True
                if italic:
                    run.italic = True
                if size:
                    run.font.size = self.Pt(size)
                if color:
                    self._set_color(run.font, color)
                if seg.kind == "placeholder":
                    run.font.highlight_color = self.HIGHLIGHT.YELLOW

    def _paragraph(self, *, style: str | None = None, before: float | None = None, after: float | None = None,
                   left: float | None = None, first: float | None = None, align=None, keep_next: bool = False):
        paragraph = self.doc.add_paragraph(style=style)
        fmt = paragraph.paragraph_format
        if before is not None:
            fmt.space_before = self.Pt(before)
        if after is not None:
            fmt.space_after = self.Pt(after)
        if left is not None:
            fmt.left_indent = self.Cm(left)
        if first is not None:
            fmt.first_line_indent = self.Cm(first)
        if align is not None:
            paragraph.alignment = align
        if keep_next:
            fmt.keep_with_next = True
        self.last_was_table = False
        return paragraph

    # -- blocks ----------------------------------------------------------------

    def title_block(self, title: str) -> None:
        paragraph = self._paragraph(style="Title")
        self._runs(paragraph, strip_inline(title), links=False)
        if self.meta:
            meta = self._paragraph(align=self.ALIGN.CENTER, after=10)
            self._runs(meta, self.meta, size=9, color=MUTED, links=False)
        self._blind_notice()

    def _blind_notice(self) -> None:
        exposed = exposed_names(self.draft, self.profile)
        if not exposed:
            return
        names = ", ".join(exposed)
        text = (f"블라인드 확인 필요 — 팀원 실명 {len(exposed)}개가 본문에 보여요: {names}. "
                "사업계획서 제출본에는 실명을 쓸 수 없으니 ○○로 가려 주세요.")
        self._box(text, fill="FDECEA", border_color="D92D20", border_val="single", color="912018", bold=True)

    def _box(self, text: str, *, fill: str, border_color: str, border_val: str, color: str,
             bold: bool = False, italic: bool = False, align=None) -> None:
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = self.TABLE_ALIGN.CENTER
        table.autofit = False
        self._table_props(table, [TEXT_WIDTH_CM], border_color=border_color, border_val=border_val, border_sz="8")
        cell = table.cell(0, 0)
        cell.width = self.Cm(TEXT_WIDTH_CM)
        self._shade(cell, fill)
        paragraph = cell.paragraphs[0]
        fmt = paragraph.paragraph_format
        fmt.space_before, fmt.space_after, fmt.line_spacing = self.Pt(4), self.Pt(4), 1.25
        if align is not None:
            paragraph.alignment = align
        self._runs(paragraph, text, size=9.5, color=color, bold=bold, italic=italic)
        self._after_table()

    def heading(self, block: Block, level_offset: int) -> None:
        text = strip_inline(block.text)
        level = max(1, min(3, block.level - level_offset))
        paragraph = self._paragraph(style=f"Heading {level}")
        self._runs(paragraph, text, links=False)
        compact = text.replace(" ", "")
        if any(h.replace(" ", "") in compact for h in _REFERENCE_HEADINGS):
            self.in_references, self.references_level = True, block.level
        elif self.in_references and block.level <= self.references_level:
            self.in_references = False

    def para(self, block: Block) -> None:
        is_sources = all(line.startswith("출처") for line in block.lines)
        paragraph = self._paragraph(after=6 if not is_sources else 2)
        size = 9 if is_sources else None
        color = NOTE if is_sources else None
        for index, line in enumerate(block.lines):
            if index:
                paragraph.add_run().add_break()
            self._runs(paragraph, line, size=size, color=color)

    def item(self, block: Block) -> None:
        marker = block.marker
        text = block.text
        is_note = marker == "※" or text.startswith("※")
        if self.in_references:
            depth = block.level
        elif marker in TOP_LEVEL_MARKERS or marker == "※" or marker[:1].isdigit():
            depth = block.level
        else:
            depth = block.level + 1
        hanging = 0.55
        left = 0.6 * depth + hanging
        size = 9 if self.in_references else 9.5 if is_note else None
        color = NOTE if is_note else None
        before = 6 if marker in TOP_LEVEL_MARKERS and not self.in_references else 1
        paragraph = self._paragraph(left=left, first=-hanging, before=before, after=2)
        display = "•" if marker == "*" else marker
        marker_run = paragraph.add_run(display + " ")
        if size:
            marker_run.font.size = self.Pt(size)
        if color:
            self._set_color(marker_run.font, color)
        if marker in TOP_LEVEL_MARKERS:
            marker_run.bold = True
            label = _LABEL_PAREN.match(text)
            if label:
                self._runs(paragraph, label.group(1), bold=True, size=size, color=color, links=False)
                text = text[label.end():]
                paragraph.add_run(" ")
        elif self.draft.channel == "instagram":
            label = _LABEL_COLON.match(text)
            if label:
                self._runs(paragraph, label.group(1) + label.group(2), bold=True, size=size, links=False)
                text = text[label.end() - 1:]
        for index, line in enumerate(text.split("\n")):
            if index:
                paragraph.add_run().add_break()
            self._runs(paragraph, line, size=size, color=color)

    def caption(self, block: Block) -> None:
        paragraph = self._paragraph(align=self.ALIGN.CENTER, before=8, after=3, keep_next=True)
        self._runs(paragraph, f"< {strip_inline(block.text)} >", bold=True, size=10, links=False)

    def quote(self, block: Block) -> None:
        paragraph = self._paragraph(left=0.6, after=4)
        self._runs(paragraph, block.text, italic=True, color=MUTED)

    def image(self, block: Block) -> None:
        self._box(f"[이미지 자리] {block.text}", fill="F3F4F6", border_color="9CA3AF", border_val="dashed",
                  color=NOTE, italic=True, align=self.ALIGN.CENTER)

    def _shade(self, cell, fill: str) -> None:
        tcpr = cell._tc.get_or_add_tcPr()
        self._insert(tcpr, self._el("w:shd", val="clear", color="auto", fill=fill))

    def _table_props(self, table, widths: list[float], *, border_color: str = BORDER, border_val: str = "single",
                     border_sz: str = "4") -> None:
        tblpr = table._tbl.tblPr
        total = int(round(sum(widths) / 2.54 * 1440))
        tblw = tblpr.find(self.qn("w:tblW"))
        if tblw is None:
            tblw = self._el("w:tblW")
            self._insert(tblpr, tblw)
        tblw.set(self.qn("w:w"), str(total))
        tblw.set(self.qn("w:type"), "dxa")
        borders = self._el("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            borders.append(self._el(f"w:{edge}", val=border_val, sz=border_sz, space="0", color=border_color))
        self._insert(tblpr, borders)
        margins = self._el("w:tblCellMar")
        for edge, value in (("top", "45"), ("left", "100"), ("bottom", "45"), ("right", "100")):
            margins.append(self._el(f"w:{edge}", w=value, type="dxa"))
        self._insert(tblpr, margins)
        for index, width in enumerate(widths):
            table.columns[index].width = self.Cm(width)

    def _after_table(self) -> None:
        spacer = self.doc.add_paragraph()
        fmt = spacer.paragraph_format
        fmt.space_before, fmt.space_after, fmt.line_spacing = self.Pt(0), self.Pt(4), 1.0
        self.last_was_table = True

    def table(self, block: Block) -> None:
        rows = [block.header, *block.rows]
        ncols = len(block.header)
        widths = column_widths(rows)
        numeric = _numeric_columns(block.rows)
        table = self.doc.add_table(rows=len(rows), cols=ncols)
        table.style = self.doc.styles["Table Grid"]
        table.alignment = self.TABLE_ALIGN.CENTER
        table.autofit = False
        self._table_props(table, widths)
        for r_index, values in enumerate(rows):
            row = table.rows[r_index]
            trpr = row._tr.get_or_add_trPr()
            self._insert(trpr, self._el("w:cantSplit"))
            if r_index == 0:
                self._insert(trpr, self._el("w:tblHeader"))
            for c_index in range(ncols):
                cell = row.cells[c_index]
                cell.width = self.Cm(widths[c_index])
                if r_index == 0:
                    self._shade(cell, HEADER_FILL)
                cell.vertical_alignment = self.VALIGN.CENTER
                paragraph = cell.paragraphs[0]
                fmt = paragraph.paragraph_format
                fmt.space_before, fmt.space_after, fmt.line_spacing = self.Pt(1), self.Pt(1), 1.15
                align = block.aligns[c_index] if c_index < len(block.aligns) else ""
                if r_index == 0 or align == "center":
                    paragraph.alignment = self.ALIGN.CENTER
                elif align == "right" or c_index in numeric:
                    paragraph.alignment = self.ALIGN.RIGHT
                self._runs(paragraph, values[c_index], size=9.5, bold=r_index == 0)
        self._after_table()

    def hashtags(self) -> None:
        tags = [t.strip() for t in self.draft.hashtags if t.strip()]
        if not tags:
            return
        lines = [ln for ln in (self.draft.content or "").strip().splitlines() if ln.strip()]
        last = lines[-1].strip() if lines else ""
        if is_hashtag_line(last) and all(t in last.split() for t in tags):
            return
        label = "태그" if self.draft.channel == "naver_blog" else "해시태그"
        paragraph = self._paragraph(before=10)
        self._runs(paragraph, f"{label}: {' '.join(tags)}", color="0B6E4F")

    # -- driver ----------------------------------------------------------------

    def build(self) -> bytes:
        self.setup()
        blocks = parse_blocks(self.draft.content or "")
        title = self.draft.title.strip()
        if blocks and blocks[0].kind == "heading" and blocks[0].level == 1:
            title = blocks[0].text.strip() or title
            blocks = blocks[1:]
        self.title_block(title or f"{self.channel_label} 초안")
        top_levels = [b.level for b in blocks if b.kind == "heading"]
        level_offset = max(0, min(top_levels) - 1) if top_levels else 0
        for block in blocks:
            if block.kind == "heading":
                self.heading(block, level_offset)
            elif block.kind == "para":
                self.para(block)
            elif block.kind == "item":
                self.item(block)
            elif block.kind == "table":
                self.table(block)
            elif block.kind == "caption":
                self.caption(block)
            elif block.kind == "quote":
                self.quote(block)
            elif block.kind == "image":
                self.image(block)
        if self.draft.channel != "bizplan":
            self.hashtags()
        buffer = io.BytesIO()
        self.doc.save(buffer)
        return buffer.getvalue()



def exposed_names(draft: Draft, profile: Profile | None) -> list[str]:
    """Team member names (from the profile) that appear in the draft — the
    business-plan blind rule. Same matching as ``channels.profile_checks``."""
    if profile is None:
        return []
    flat = re.sub(r"\s", "", f"{draft.title}\n{draft.content}").lower()
    names: list[str] = []
    for member in profile.team:
        name = (member.name or "").strip()
        key = re.sub(r"\s", "", name).lower()
        if len(name) >= 2 and key and key in flat and name not in names:
            names.append(name)
    return names


def build_docx(draft: Draft, *, meta: str = "", profile: Profile | None = None, channel_label: str = "") -> bytes:
    """Render ``draft`` as a .docx document and return the file bytes."""
    return _Builder(draft, meta=meta, profile=profile, channel_label=channel_label or draft.channel).build()
